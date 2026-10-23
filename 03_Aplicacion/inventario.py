"""Motor de lectura, validación, estructuración e inicialización segura del inventario Excel.

Proyecto: FerreteríaMS v0
Tarea: FERMS-V0-P02
Requisitos relacionados: RF-020, RF-022, RF-023, RF-024, RF-025, RF-029, RF-030, RF-031.
Criterios relacionados: CA-020 (parcial), CA-021 (parcial), CA-022, CA-047, CA-048, CA-049, CA-050.

Importar este módulo no crea ni modifica archivos en disco.
"""

from __future__ import annotations

import io
from dataclasses import dataclass, field
from enum import Enum
from pathlib import Path
from typing import Any, BinaryIO, Mapping
import unicodedata
from zipfile import BadZipFile

import openpyxl
from openpyxl.utils.exceptions import InvalidFileException
import pandas as pd

# Nombres exactos de hojas y columnas aprobados en 02_Especificaciones_v0.md (Sección 3.3)
# y extendidos en 03_Especificaciones_Final.md (Sección 3.4 — campo lógico Activo en Ubicaciones para P14)
HOJA_PRODUCTOS = "Productos"
HOJA_CATEGORIAS = "Categorías"
HOJA_UBICACIONES = "Ubicaciones"
HOJA_ASIGNACIONES = "Asignaciones"

ESQUEMA_HOJAS: dict[str, list[str]] = {
    HOJA_PRODUCTOS: ["Codigo_Producto", "Nombre", "Categoria_ID", "Descripcion", "Activo"],
    HOJA_CATEGORIAS: ["Categoria_ID", "Nombre_Categoria"],
    HOJA_UBICACIONES: ["Ubicacion_ID", "Rack", "Columna", "Nivel", "Posicion", "Activo"],
    HOJA_ASIGNACIONES: ["Codigo_Producto", "Ubicacion_ID", "Principal"],
}

COLUMNAS_OPCIONALES_RETROCOMPATIBLES: dict[str, frozenset[str]] = {
    HOJA_PRODUCTOS: frozenset({"Activo"}),
    HOJA_UBICACIONES: frozenset({"Activo"}),
}

ORDEN_HOJAS: tuple[str, ...] = (
    HOJA_PRODUCTOS,
    HOJA_CATEGORIAS,
    HOJA_UBICACIONES,
    HOJA_ASIGNACIONES,
)

VALORES_PRINCIPAL_SI: frozenset[str] = frozenset({"sí", "si"})
VALORES_PRINCIPAL_NO: frozenset[str] = frozenset({"no"})

VALORES_ACTIVO_SI: frozenset[str] = frozenset({"sí", "si", "true", "1", "activo", "activa"})
VALORES_ACTIVO_NO: frozenset[str] = frozenset({"no", "false", "0", "inactivo", "inactiva"})

MARCA_DEMOSTRATIVA = "[DEMOSTRATIVO]"


class CodigoErrorInventario(str, Enum):
    """Códigos identificables para cada tipo de error del motor de inventario."""

    ARCHIVO_INEXISTENTE = "ARCHIVO_INEXISTENTE"
    ARCHIVO_CORRUPTO = "ARCHIVO_CORRUPTO"
    HOJA_FALTANTE = "HOJA_FALTANTE"
    HOJA_NO_AUTORIZADA = "HOJA_NO_AUTORIZADA"
    COLUMNA_FALTANTE = "COLUMNA_FALTANTE"
    COLUMNA_NO_AUTORIZADA = "COLUMNA_NO_AUTORIZADA"
    VALOR_OBLIGATORIO_AUSENTE = "VALOR_OBLIGATORIO_AUSENTE"
    IDENTIFICADOR_DUPLICADO = "IDENTIFICADOR_DUPLICADO"
    NOMBRE_CATEGORIA_DUPLICADO = "NOMBRE_CATEGORIA_DUPLICADO"
    CATEGORIA_INEXISTENTE = "CATEGORIA_INEXISTENTE"
    UBICACION_INEXISTENTE = "UBICACION_INEXISTENTE"
    UBICACION_INACTIVA_ASIGNADA = "UBICACION_INACTIVA_ASIGNADA"
    PRODUCTO_INEXISTENTE = "PRODUCTO_INEXISTENTE"
    UBICACION_FISICA_DUPLICADA = "UBICACION_FISICA_DUPLICADA"
    ASIGNACION_DUPLICADA = "ASIGNACION_DUPLICADA"
    PRODUCTO_SIN_ASIGNACION = "PRODUCTO_SIN_ASIGNACION"
    UBICACION_PRINCIPAL_MULTIPLE = "UBICACION_PRINCIPAL_MULTIPLE"
    SIN_UBICACION_PRINCIPAL = "SIN_UBICACION_PRINCIPAL"
    VALOR_PRINCIPAL_INVALIDO = "VALOR_PRINCIPAL_INVALIDO"
    VALOR_ACTIVO_INVALIDO = "VALOR_ACTIVO_INVALIDO"
    PROTECCION_INVENTARIO_EXISTENTE = "PROTECCION_INVENTARIO_EXISTENTE"
    ESTADO_REMOTO_DESCONOCIDO = "ESTADO_REMOTO_DESCONOCIDO"


class EstadoExistenciaRemota(str, Enum):
    """Estados posibles al consultar la existencia del archivo maestro en GitHub."""

    EXISTE = "EXISTE"
    NO_EXISTE = "NO_EXISTE"
    DESCONOCIDO_ERROR = "DESCONOCIDO_ERROR"


@dataclass(frozen=True)
class DetalleErrorInventario:
    """Detalle estructurado de un error detectado en el inventario."""

    codigo: CodigoErrorInventario
    mensaje: str
    hoja: str | None = None
    fila: int | None = None
    campo: str | None = None
    valor: str | None = None


class ErrorInventario(Exception):
    """Excepción base del motor de inventario."""

    codigo_por_defecto: CodigoErrorInventario = CodigoErrorInventario.ARCHIVO_CORRUPTO

    def __init__(
        self,
        mensaje: str,
        codigo: CodigoErrorInventario | None = None,
        errores: list[DetalleErrorInventario] | None = None,
    ) -> None:
        super().__init__(mensaje)
        self.mensaje = mensaje
        self.codigo = codigo or self.codigo_por_defecto
        self.errores: list[DetalleErrorInventario] = list(errores) if errores else [
            DetalleErrorInventario(codigo=self.codigo, mensaje=mensaje)
        ]


class ErrorArchivoInexistente(ErrorInventario):
    """El archivo Excel indicado no existe."""

    codigo_por_defecto = CodigoErrorInventario.ARCHIVO_INEXISTENTE


class ErrorArchivoCorrupto(ErrorInventario):
    """El archivo Excel está corrupto, vacío o es ilegible."""

    codigo_por_defecto = CodigoErrorInventario.ARCHIVO_CORRUPTO


class ErrorEstructuraInventario(ErrorInventario):
    """La estructura de hojas o columnas del archivo Excel es inválida."""


class ErrorHojaFaltante(ErrorEstructuraInventario):
    """Falta al menos una de las cuatro hojas obligatorias o existen hojas no autorizadas."""

    codigo_por_defecto = CodigoErrorInventario.HOJA_FALTANTE


class ErrorHojaNoAutorizada(ErrorHojaFaltante):
    """Existen hojas adicionales no autorizadas en el archivo Excel."""

    codigo_por_defecto = CodigoErrorInventario.HOJA_NO_AUTORIZADA


class ErrorColumnaFaltante(ErrorEstructuraInventario):
    """Faltan columnas obligatorias o existen columnas no autorizadas."""

    codigo_por_defecto = CodigoErrorInventario.COLUMNA_FALTANTE


class ErrorValidacionInventario(ErrorInventario):
    """Los datos del inventario incumplen las reglas de integridad."""


class ErrorValorObligatorio(ErrorValidacionInventario):
    """Falta un valor obligatorio en alguna fila."""

    codigo_por_defecto = CodigoErrorInventario.VALOR_OBLIGATORIO_AUSENTE


class ErrorIdentificadorDuplicado(ErrorValidacionInventario):
    """Existe un identificador único duplicado (producto, categoría o ubicación)."""

    codigo_por_defecto = CodigoErrorInventario.IDENTIFICADOR_DUPLICADO


class ErrorCategoriaInexistente(ErrorValidacionInventario):
    """Un producto referencia una categoría que no existe en la hoja Categorías."""

    codigo_por_defecto = CodigoErrorInventario.CATEGORIA_INEXISTENTE


class ErrorUbicacionInexistente(ErrorValidacionInventario):
    """Una asignación referencia una ubicación que no existe en la hoja Ubicaciones."""

    codigo_por_defecto = CodigoErrorInventario.UBICACION_INEXISTENTE


class ErrorProductoInexistente(ErrorValidacionInventario):
    """Una asignación referencia un código de producto que no existe en Productos."""

    codigo_por_defecto = CodigoErrorInventario.PRODUCTO_INEXISTENTE


class ErrorUbicacionFisicaDuplicada(ErrorValidacionInventario):
    """Dos ubicaciones tienen la misma combinación física Rack + Columna + Nivel + Posicion."""

    codigo_por_defecto = CodigoErrorInventario.UBICACION_FISICA_DUPLICADA


class ErrorAsignacionDuplicada(ErrorValidacionInventario):
    """Se repite la relación entre el mismo producto y la misma ubicación."""

    codigo_por_defecto = CodigoErrorInventario.ASIGNACION_DUPLICADA


class ErrorProductoSinAsignacion(ErrorValidacionInventario):
    """Un producto registrado no posee ninguna ubicación asignada."""

    codigo_por_defecto = CodigoErrorInventario.PRODUCTO_SIN_ASIGNACION


class ErrorUbicacionPrincipalMultiple(ErrorValidacionInventario):
    """Un producto tiene más de una ubicación marcada como principal."""

    codigo_por_defecto = CodigoErrorInventario.UBICACION_PRINCIPAL_MULTIPLE


class ErrorSinUbicacionPrincipal(ErrorValidacionInventario):
    """Un producto tiene asignaciones, pero ninguna marcada como ubicación principal."""

    codigo_por_defecto = CodigoErrorInventario.SIN_UBICACION_PRINCIPAL


class ErrorValorPrincipalInvalido(ErrorValidacionInventario):
    """El campo Principal contiene un valor ambiguo o distinto de Sí/No."""

    codigo_por_defecto = CodigoErrorInventario.VALOR_PRINCIPAL_INVALIDO


class ErrorNombreCategoriaDuplicado(ErrorValidacionInventario):
    """Dos categorías tienen el mismo nombre normalizado."""

    codigo_por_defecto = CodigoErrorInventario.NOMBRE_CATEGORIA_DUPLICADO


class ErrorUbicacionInactivaAsignada(ErrorValidacionInventario):
    """Una asignación referencia una ubicación que se encuentra desactivada."""

    codigo_por_defecto = CodigoErrorInventario.UBICACION_INACTIVA_ASIGNADA


class ErrorValorActivoInvalido(ErrorValidacionInventario):
    """El campo Activo contiene un valor ambiguo o distinto de Sí/No."""

    codigo_por_defecto = CodigoErrorInventario.VALOR_ACTIVO_INVALIDO


class ErrorProteccionInventario(ErrorInventario):
    """Operación bloqueada para proteger un inventario existente o ante estado incierto."""

    codigo_por_defecto = CodigoErrorInventario.PROTECCION_INVENTARIO_EXISTENTE


MAPA_EXCEPCIONES: dict[CodigoErrorInventario, type[ErrorInventario]] = {
    CodigoErrorInventario.ARCHIVO_INEXISTENTE: ErrorArchivoInexistente,
    CodigoErrorInventario.ARCHIVO_CORRUPTO: ErrorArchivoCorrupto,
    CodigoErrorInventario.HOJA_FALTANTE: ErrorHojaFaltante,
    CodigoErrorInventario.HOJA_NO_AUTORIZADA: ErrorHojaNoAutorizada,
    CodigoErrorInventario.COLUMNA_FALTANTE: ErrorColumnaFaltante,
    CodigoErrorInventario.COLUMNA_NO_AUTORIZADA: ErrorColumnaFaltante,
    CodigoErrorInventario.VALOR_OBLIGATORIO_AUSENTE: ErrorValorObligatorio,
    CodigoErrorInventario.IDENTIFICADOR_DUPLICADO: ErrorIdentificadorDuplicado,
    CodigoErrorInventario.NOMBRE_CATEGORIA_DUPLICADO: ErrorNombreCategoriaDuplicado,
    CodigoErrorInventario.CATEGORIA_INEXISTENTE: ErrorCategoriaInexistente,
    CodigoErrorInventario.UBICACION_INEXISTENTE: ErrorUbicacionInexistente,
    CodigoErrorInventario.UBICACION_INACTIVA_ASIGNADA: ErrorUbicacionInactivaAsignada,
    CodigoErrorInventario.PRODUCTO_INEXISTENTE: ErrorProductoInexistente,
    CodigoErrorInventario.UBICACION_FISICA_DUPLICADA: ErrorUbicacionFisicaDuplicada,
    CodigoErrorInventario.ASIGNACION_DUPLICADA: ErrorAsignacionDuplicada,
    CodigoErrorInventario.PRODUCTO_SIN_ASIGNACION: ErrorProductoSinAsignacion,
    CodigoErrorInventario.UBICACION_PRINCIPAL_MULTIPLE: ErrorUbicacionPrincipalMultiple,
    CodigoErrorInventario.SIN_UBICACION_PRINCIPAL: ErrorSinUbicacionPrincipal,
    CodigoErrorInventario.VALOR_PRINCIPAL_INVALIDO: ErrorValorPrincipalInvalido,
    CodigoErrorInventario.VALOR_ACTIVO_INVALIDO: ErrorValorActivoInvalido,
    CodigoErrorInventario.PROTECCION_INVENTARIO_EXISTENTE: ErrorProteccionInventario,
    CodigoErrorInventario.ESTADO_REMOTO_DESCONOCIDO: ErrorProteccionInventario,
}


@dataclass
class ResultadoValidacion:
    """Resultado de la validación integral de las cuatro hojas del inventario."""

    es_valido: bool
    errores: list[DetalleErrorInventario] = field(default_factory=list)

    @property
    def codigos_error(self) -> list[CodigoErrorInventario]:
        return [e.codigo for e in self.errores]


@dataclass
class InventarioDatos:
    """Contenedor validado de las cuatro hojas del inventario con datos normalizados como texto."""

    productos: pd.DataFrame
    categorias: pd.DataFrame
    ubicaciones: pd.DataFrame
    asignaciones: pd.DataFrame

    def obtener_hojas(self) -> dict[str, pd.DataFrame]:
        """Devuelve copias de las cuatro hojas en el orden oficial aprobado."""
        return {
            HOJA_PRODUCTOS: self.productos.copy(),
            HOJA_CATEGORIAS: self.categorias.copy(),
            HOJA_UBICACIONES: self.ubicaciones.copy(),
            HOJA_ASIGNACIONES: self.asignaciones.copy(),
        }

    def obtener_productos_estructurados(
        self, *, incluir_inactivos: bool = True
    ) -> list[dict[str, Any]]:
        """Prepara la vista estructurada de productos para las funcionalidades de búsqueda y administración."""
        return obtener_productos_estructurados(self, incluir_inactivos=incluir_inactivos)

    def obtener_productos_activos_estructurados(self) -> list[dict[str, Any]]:
        """Devuelve únicamente los productos activos con su estructura completa."""
        return obtener_productos_activos_estructurados(self)

    def obtener_producto_por_codigo(
        self, codigo_producto: str, *, incluir_inactivos: bool = True
    ) -> dict[str, Any] | None:
        """Devuelve la estructura detallada de un producto por su código exacto."""
        return obtener_producto_por_codigo(
            self, codigo_producto, incluir_inactivos=incluir_inactivos
        )

    def construir_tabla_consulta(self) -> pd.DataFrame:
        """Devuelve un DataFrame relacional consolidado de productos, categorías y ubicaciones."""
        return construir_tabla_consulta(self)


def _limpiar_texto(valor: Any) -> str:
    """Convierte cualquier valor escalar a cadena limpia preservando ceros iniciales."""
    if valor is None or pd.isna(valor):
        return ""
    return str(valor).strip()


def normalizar_texto_comparacion_catalogo(valor: Any) -> str:
    """Normaliza un texto para comparación de unicidad (sin acentos, minúsculas y espacios simples)."""
    texto = " ".join(_limpiar_texto(valor).split())
    if not texto:
        return ""
    descompuesto = unicodedata.normalize("NFD", texto)
    sin_acentos = "".join(c for c in descompuesto if unicodedata.category(c) != "Mn")
    return unicodedata.normalize("NFC", sin_acentos).casefold()


def normalizar_valor_principal(valor: Any) -> str:
    """Normaliza el campo Principal a 'Sí' o 'No'.

    Rechaza cualquier valor vacío, ambiguo o distinto de las representaciones de Sí/No.
    """
    texto = _limpiar_texto(valor)
    texto_min = texto.lower()
    if texto_min in VALORES_PRINCIPAL_SI:
        return "Sí"
    if texto_min in VALORES_PRINCIPAL_NO:
        return "No"
    raise ErrorValorPrincipalInvalido(
        f"Valor inválido para el campo 'Principal': {texto!r}. Solo se admiten 'Sí' o 'No'.",
        codigo=CodigoErrorInventario.VALOR_PRINCIPAL_INVALIDO,
    )


def normalizar_valor_activo(valor: Any, *, permitir_vacio_retrocompatible: bool = True) -> str:
    """Normaliza el campo lógico Activo a 'Sí' o 'No'.

    Por compatibilidad retroactiva con archivos v0 (CA-F-009), un valor ausente o vacío
    se interpreta como activo ('Sí') cuando `permitir_vacio_retrocompatible=True`.
    """
    texto = _limpiar_texto(valor)
    if not texto and permitir_vacio_retrocompatible:
        return "Sí"
    texto_min = texto.lower()
    if texto_min in VALORES_ACTIVO_SI:
        return "Sí"
    if texto_min in VALORES_ACTIVO_NO:
        return "No"
    raise ErrorValorActivoInvalido(
        f"Valor inválido para el campo 'Activo': {texto!r}. Solo se admiten 'Sí' o 'No'.",
        codigo=CodigoErrorInventario.VALOR_ACTIVO_INVALIDO,
    )


def _normalizar_dataframe_texto(df: pd.DataFrame, columnas_esperadas: list[str]) -> pd.DataFrame:
    """Normaliza un DataFrame conservando todas las celdas como cadenas de texto limpias."""
    copia = df.copy()
    for col in columnas_esperadas:
        if col in copia.columns:
            copia[col] = copia[col].map(_limpiar_texto).astype(str)
            if col == "Activo":
                copia[col] = copia[col].map(lambda v: "Sí" if not v else v).astype(str)
        elif col == "Activo":
            copia[col] = "Sí"
    return copia[columnas_esperadas].reset_index(drop=True)


def _lanzar_excepcion_desde_errores(errores: list[DetalleErrorInventario]) -> None:
    """Construye y lanza la excepción específica correspondiente al primer error registrado."""
    if not errores:
        return
    primer_error = errores[0]
    clase_exc = MAPA_EXCEPCIONES.get(primer_error.codigo, ErrorValidacionInventario)
    resumen = "; ".join(f"[{e.codigo.value}] {e.mensaje}" for e in errores)
    raise clase_exc(resumen, codigo=primer_error.codigo, errores=errores)


def validar_estructura_hojas(
    hojas: Mapping[str, pd.DataFrame],
    lanzar_excepcion: bool = True,
) -> ResultadoValidacion:
    """Verifica que existan exactamente las cuatro hojas aprobadas y sus columnas exactas."""
    errores: list[DetalleErrorInventario] = []

    if not isinstance(hojas, Mapping) or len(hojas) == 0:
        errores.append(
            DetalleErrorInventario(
                codigo=CodigoErrorInventario.HOJA_FALTANTE,
                mensaje="El libro no contiene hojas válidas.",
            )
        )
        if lanzar_excepcion:
            _lanzar_excepcion_desde_errores(errores)
        return ResultadoValidacion(es_valido=False, errores=errores)

    hojas_no_autorizadas = [str(nombre) for nombre in hojas.keys() if nombre not in ESQUEMA_HOJAS]
    if hojas_no_autorizadas:
        errores.append(
            DetalleErrorInventario(
                codigo=CodigoErrorInventario.HOJA_NO_AUTORIZADA,
                mensaje=(
                    "El archivo contiene hojas adicionales no autorizadas para v0: "
                    f"{', '.join(hojas_no_autorizadas)}. "
                    f"Solo se admiten exactamente las cuatro hojas aprobadas: {', '.join(ORDEN_HOJAS)}."
                ),
                hoja=", ".join(hojas_no_autorizadas),
            )
        )

    for nombre_hoja, columnas_requeridas in ESQUEMA_HOJAS.items():
        if nombre_hoja not in hojas:
            errores.append(
                DetalleErrorInventario(
                    codigo=CodigoErrorInventario.HOJA_FALTANTE,
                    mensaje=f"Falta la hoja obligatoria '{nombre_hoja}'.",
                    hoja=nombre_hoja,
                )
            )
            continue

        df = hojas[nombre_hoja]
        if not isinstance(df, pd.DataFrame):
            errores.append(
                DetalleErrorInventario(
                    codigo=CodigoErrorInventario.ARCHIVO_CORRUPTO,
                    mensaje=f"El contenido de la hoja '{nombre_hoja}' no es una tabla válida.",
                    hoja=nombre_hoja,
                )
            )
            continue

        columnas_actuales = [str(c) for c in df.columns]
        columnas_opcionales = COLUMNAS_OPCIONALES_RETROCOMPATIBLES.get(nombre_hoja, frozenset())
        faltantes = [
            c for c in columnas_requeridas
            if c not in columnas_actuales and c not in columnas_opcionales
        ]
        if faltantes:
            errores.append(
                DetalleErrorInventario(
                    codigo=CodigoErrorInventario.COLUMNA_FALTANTE,
                    mensaje=(
                        f"En la hoja '{nombre_hoja}' faltan las columnas obligatorias: "
                        f"{', '.join(faltantes)}."
                    ),
                    hoja=nombre_hoja,
                    campo=", ".join(faltantes),
                )
            )

        no_autorizadas = [c for c in columnas_actuales if c not in columnas_requeridas]
        if no_autorizadas:
            errores.append(
                DetalleErrorInventario(
                    codigo=CodigoErrorInventario.COLUMNA_NO_AUTORIZADA,
                    mensaje=(
                        f"En la hoja '{nombre_hoja}' existen columnas no autorizadas para v0: "
                        f"{', '.join(no_autorizadas)}."
                    ),
                    hoja=nombre_hoja,
                    campo=", ".join(no_autorizadas),
                )
            )

    if errores and lanzar_excepcion:
        _lanzar_excepcion_desde_errores(errores)

    return ResultadoValidacion(es_valido=len(errores) == 0, errores=errores)


def validar_inventario(
    hojas: Mapping[str, pd.DataFrame],
    lanzar_excepcion: bool = True,
) -> ResultadoValidacion:
    """Valida integralmente la estructura, los tipos, las claves únicas y las relaciones."""
    resultado_estructura = validar_estructura_hojas(hojas, lanzar_excepcion=lanzar_excepcion)
    if not resultado_estructura.es_valido:
        return resultado_estructura

    errores: list[DetalleErrorInventario] = []

    df_productos = _normalizar_dataframe_texto(
        hojas[HOJA_PRODUCTOS], ESQUEMA_HOJAS[HOJA_PRODUCTOS]
    )
    df_categorias = _normalizar_dataframe_texto(
        hojas[HOJA_CATEGORIAS], ESQUEMA_HOJAS[HOJA_CATEGORIAS]
    )
    df_ubicaciones = _normalizar_dataframe_texto(
        hojas[HOJA_UBICACIONES], ESQUEMA_HOJAS[HOJA_UBICACIONES]
    )
    df_asignaciones = _normalizar_dataframe_texto(
        hojas[HOJA_ASIGNACIONES], ESQUEMA_HOJAS[HOJA_ASIGNACIONES]
    )

    # 1. Validación de Categorías
    categorias_vistas: set[str] = set()
    nombres_categorias_vistos: set[str] = set()
    for idx, fila in df_categorias.iterrows():
        num_fila = int(idx) + 2
        cat_id = fila["Categoria_ID"]
        nombre_cat = fila["Nombre_Categoria"]

        if not cat_id:
            errores.append(
                DetalleErrorInventario(
                    codigo=CodigoErrorInventario.VALOR_OBLIGATORIO_AUSENTE,
                    mensaje=f"Fila {num_fila} en '{HOJA_CATEGORIAS}': 'Categoria_ID' es obligatorio.",
                    hoja=HOJA_CATEGORIAS,
                    fila=num_fila,
                    campo="Categoria_ID",
                )
            )
        elif cat_id in categorias_vistas:
            errores.append(
                DetalleErrorInventario(
                    codigo=CodigoErrorInventario.IDENTIFICADOR_DUPLICADO,
                    mensaje=f"Fila {num_fila} en '{HOJA_CATEGORIAS}': 'Categoria_ID' duplicado ('{cat_id}').",
                    hoja=HOJA_CATEGORIAS,
                    fila=num_fila,
                    campo="Categoria_ID",
                    valor=cat_id,
                )
            )
        else:
            categorias_vistas.add(cat_id)

        if not nombre_cat:
            errores.append(
                DetalleErrorInventario(
                    codigo=CodigoErrorInventario.VALOR_OBLIGATORIO_AUSENTE,
                    mensaje=f"Fila {num_fila} en '{HOJA_CATEGORIAS}': 'Nombre_Categoria' es obligatorio.",
                    hoja=HOJA_CATEGORIAS,
                    fila=num_fila,
                    campo="Nombre_Categoria",
                )
            )
        else:
            clave_nombre_cat = normalizar_texto_comparacion_catalogo(nombre_cat)
            if clave_nombre_cat in nombres_categorias_vistos:
                errores.append(
                    DetalleErrorInventario(
                        codigo=CodigoErrorInventario.NOMBRE_CATEGORIA_DUPLICADO,
                        mensaje=(
                            f"Fila {num_fila} en '{HOJA_CATEGORIAS}': "
                            f"'Nombre_Categoria' duplicado ('{nombre_cat}')."
                        ),
                        hoja=HOJA_CATEGORIAS,
                        fila=num_fila,
                        campo="Nombre_Categoria",
                        valor=nombre_cat,
                    )
                )
            else:
                nombres_categorias_vistos.add(clave_nombre_cat)

    # 2. Validación de Ubicaciones
    ubicaciones_vistas: set[str] = set()
    ubicaciones_inactivas: set[str] = set()
    combinaciones_fisicas_vistas: set[tuple[str, str, str, str]] = set()
    for idx, fila in df_ubicaciones.iterrows():
        num_fila = int(idx) + 2
        ubic_id = fila["Ubicacion_ID"]
        rack = fila["Rack"]
        columna = fila["Columna"]
        nivel = fila["Nivel"]
        posicion = fila["Posicion"]
        val_activo = fila.get("Activo", "Sí")

        if not ubic_id:
            errores.append(
                DetalleErrorInventario(
                    codigo=CodigoErrorInventario.VALOR_OBLIGATORIO_AUSENTE,
                    mensaje=f"Fila {num_fila} en '{HOJA_UBICACIONES}': 'Ubicacion_ID' es obligatorio.",
                    hoja=HOJA_UBICACIONES,
                    fila=num_fila,
                    campo="Ubicacion_ID",
                )
            )
        elif ubic_id in ubicaciones_vistas:
            errores.append(
                DetalleErrorInventario(
                    codigo=CodigoErrorInventario.IDENTIFICADOR_DUPLICADO,
                    mensaje=f"Fila {num_fila} en '{HOJA_UBICACIONES}': 'Ubicacion_ID' duplicado ('{ubic_id}').",
                    hoja=HOJA_UBICACIONES,
                    fila=num_fila,
                    campo="Ubicacion_ID",
                    valor=ubic_id,
                )
            )
        else:
            ubicaciones_vistas.add(ubic_id)

        try:
            activo_norm = normalizar_valor_activo(val_activo, permitir_vacio_retrocompatible=True)
            if activo_norm == "No" and ubic_id:
                ubicaciones_inactivas.add(ubic_id)
        except ErrorValorActivoInvalido:
            errores.append(
                DetalleErrorInventario(
                    codigo=CodigoErrorInventario.VALOR_ACTIVO_INVALIDO,
                    mensaje=(
                        f"Fila {num_fila} en '{HOJA_UBICACIONES}': valor inválido '{val_activo}' "
                        f"en 'Activo'. Solo se admiten los estados lógicos 'Sí' y 'No'."
                    ),
                    hoja=HOJA_UBICACIONES,
                    fila=num_fila,
                    campo="Activo",
                    valor=str(val_activo),
                )
            )

        campos_fisicos = {"Rack": rack, "Columna": columna, "Nivel": nivel, "Posicion": posicion}
        faltan_fisicos = False
        for nombre_campo, valor_campo in campos_fisicos.items():
            if not valor_campo:
                faltan_fisicos = True
                errores.append(
                    DetalleErrorInventario(
                        codigo=CodigoErrorInventario.VALOR_OBLIGATORIO_AUSENTE,
                        mensaje=(
                            f"Fila {num_fila} en '{HOJA_UBICACIONES}': "
                            f"el campo '{nombre_campo}' es obligatorio."
                        ),
                        hoja=HOJA_UBICACIONES,
                        fila=num_fila,
                        campo=nombre_campo,
                    )
                )

        if not faltan_fisicos:
            clave_fisica = (
                normalizar_texto_comparacion_catalogo(rack),
                normalizar_texto_comparacion_catalogo(columna),
                normalizar_texto_comparacion_catalogo(nivel),
                normalizar_texto_comparacion_catalogo(posicion),
            )
            if clave_fisica in combinaciones_fisicas_vistas:
                etiqueta = f"{rack}-{columna}-{nivel}-{posicion}"
                errores.append(
                    DetalleErrorInventario(
                        codigo=CodigoErrorInventario.UBICACION_FISICA_DUPLICADA,
                        mensaje=(
                            f"Fila {num_fila} en '{HOJA_UBICACIONES}': combinación física "
                            f"Rack/Columna/Nivel/Posicion duplicada ('{etiqueta}')."
                        ),
                        hoja=HOJA_UBICACIONES,
                        fila=num_fila,
                        campo="Rack,Columna,Nivel,Posicion",
                        valor=etiqueta,
                    )
                )
            else:
                combinaciones_fisicas_vistas.add(clave_fisica)

    # 3. Validación de Productos
    productos_vistos: set[str] = set()
    orden_productos: list[str] = []
    for idx, fila in df_productos.iterrows():
        num_fila = int(idx) + 2
        cod_prod = fila["Codigo_Producto"]
        nombre_prod = fila["Nombre"]
        cat_id = fila["Categoria_ID"]
        val_activo_prod = fila.get("Activo", "Sí")

        if not cod_prod:
            errores.append(
                DetalleErrorInventario(
                    codigo=CodigoErrorInventario.VALOR_OBLIGATORIO_AUSENTE,
                    mensaje=f"Fila {num_fila} en '{HOJA_PRODUCTOS}': 'Codigo_Producto' es obligatorio.",
                    hoja=HOJA_PRODUCTOS,
                    fila=num_fila,
                    campo="Codigo_Producto",
                )
            )
        elif cod_prod in productos_vistos:
            errores.append(
                DetalleErrorInventario(
                    codigo=CodigoErrorInventario.IDENTIFICADOR_DUPLICADO,
                    mensaje=f"Fila {num_fila} en '{HOJA_PRODUCTOS}': 'Codigo_Producto' duplicado ('{cod_prod}').",
                    hoja=HOJA_PRODUCTOS,
                    fila=num_fila,
                    campo="Codigo_Producto",
                    valor=cod_prod,
                )
            )
        else:
            productos_vistos.add(cod_prod)
            orden_productos.append(cod_prod)

        if not nombre_prod:
            errores.append(
                DetalleErrorInventario(
                    codigo=CodigoErrorInventario.VALOR_OBLIGATORIO_AUSENTE,
                    mensaje=f"Fila {num_fila} en '{HOJA_PRODUCTOS}': 'Nombre' es obligatorio.",
                    hoja=HOJA_PRODUCTOS,
                    fila=num_fila,
                    campo="Nombre",
                )
            )

        if not cat_id:
            errores.append(
                DetalleErrorInventario(
                    codigo=CodigoErrorInventario.VALOR_OBLIGATORIO_AUSENTE,
                    mensaje=f"Fila {num_fila} en '{HOJA_PRODUCTOS}': 'Categoria_ID' es obligatorio.",
                    hoja=HOJA_PRODUCTOS,
                    fila=num_fila,
                    campo="Categoria_ID",
                )
            )
        elif cat_id not in categorias_vistas:
            errores.append(
                DetalleErrorInventario(
                    codigo=CodigoErrorInventario.CATEGORIA_INEXISTENTE,
                    mensaje=(
                        f"Fila {num_fila} en '{HOJA_PRODUCTOS}': la categoría '{cat_id}' "
                        f"asignada al producto '{cod_prod}' no existe en '{HOJA_CATEGORIAS}'."
                    ),
                    hoja=HOJA_PRODUCTOS,
                    fila=num_fila,
                    campo="Categoria_ID",
                    valor=cat_id,
                )
            )

        try:
            normalizar_valor_activo(val_activo_prod)
        except ErrorValorActivoInvalido:
            errores.append(
                DetalleErrorInventario(
                    codigo=CodigoErrorInventario.VALOR_ACTIVO_INVALIDO,
                    mensaje=(
                        f"Fila {num_fila} en '{HOJA_PRODUCTOS}': valor inválido '{val_activo_prod}' "
                        f"en 'Activo'. Solo se admiten los estados lógicos 'Sí' y 'No'."
                    ),
                    hoja=HOJA_PRODUCTOS,
                    fila=num_fila,
                    campo="Activo",
                    valor=val_activo_prod,
                )
            )

    # 4. Validación de Asignaciones
    pares_asignacion_vistos: set[tuple[str, str]] = set()
    conteo_asignaciones_por_producto: dict[str, int] = {cod: 0 for cod in orden_productos}
    conteo_principales_por_producto: dict[str, int] = {cod: 0 for cod in orden_productos}

    for idx, fila in df_asignaciones.iterrows():
        num_fila = int(idx) + 2
        cod_prod = fila["Codigo_Producto"]
        ubic_id = fila["Ubicacion_ID"]
        val_principal = fila["Principal"]

        prod_valido = False
        ubic_valida = False

        if not cod_prod:
            errores.append(
                DetalleErrorInventario(
                    codigo=CodigoErrorInventario.VALOR_OBLIGATORIO_AUSENTE,
                    mensaje=f"Fila {num_fila} en '{HOJA_ASIGNACIONES}': 'Codigo_Producto' es obligatorio.",
                    hoja=HOJA_ASIGNACIONES,
                    fila=num_fila,
                    campo="Codigo_Producto",
                )
            )
        elif cod_prod not in productos_vistos:
            errores.append(
                DetalleErrorInventario(
                    codigo=CodigoErrorInventario.PRODUCTO_INEXISTENTE,
                    mensaje=(
                        f"Fila {num_fila} en '{HOJA_ASIGNACIONES}': el producto '{cod_prod}' "
                        f"no existe en '{HOJA_PRODUCTOS}'."
                    ),
                    hoja=HOJA_ASIGNACIONES,
                    fila=num_fila,
                    campo="Codigo_Producto",
                    valor=cod_prod,
                )
            )
        else:
            prod_valido = True

        if not ubic_id:
            errores.append(
                DetalleErrorInventario(
                    codigo=CodigoErrorInventario.VALOR_OBLIGATORIO_AUSENTE,
                    mensaje=f"Fila {num_fila} en '{HOJA_ASIGNACIONES}': 'Ubicacion_ID' es obligatorio.",
                    hoja=HOJA_ASIGNACIONES,
                    fila=num_fila,
                    campo="Ubicacion_ID",
                )
            )
        elif ubic_id not in ubicaciones_vistas:
            errores.append(
                DetalleErrorInventario(
                    codigo=CodigoErrorInventario.UBICACION_INEXISTENTE,
                    mensaje=(
                        f"Fila {num_fila} en '{HOJA_ASIGNACIONES}': la ubicación '{ubic_id}' "
                        f"no existe en '{HOJA_UBICACIONES}'."
                    ),
                    hoja=HOJA_ASIGNACIONES,
                    fila=num_fila,
                    campo="Ubicacion_ID",
                    valor=ubic_id,
                )
            )
        elif ubic_id in ubicaciones_inactivas:
            errores.append(
                DetalleErrorInventario(
                    codigo=CodigoErrorInventario.UBICACION_INACTIVA_ASIGNADA,
                    mensaje=(
                        f"Fila {num_fila} en '{HOJA_ASIGNACIONES}': la ubicación '{ubic_id}' "
                        f"se encuentra inactiva y no puede estar asignada al producto '{cod_prod}'."
                    ),
                    hoja=HOJA_ASIGNACIONES,
                    fila=num_fila,
                    campo="Ubicacion_ID",
                    valor=ubic_id,
                )
            )
        else:
            ubic_valida = True

        principal_normalizado: str | None = None
        try:
            principal_normalizado = normalizar_valor_principal(val_principal)
        except ErrorValorPrincipalInvalido:
            errores.append(
                DetalleErrorInventario(
                    codigo=CodigoErrorInventario.VALOR_PRINCIPAL_INVALIDO,
                    mensaje=(
                        f"Fila {num_fila} en '{HOJA_ASIGNACIONES}': valor inválido '{val_principal}' "
                        f"en 'Principal'. Solo se admiten los estados lógicos 'Sí' y 'No'."
                    ),
                    hoja=HOJA_ASIGNACIONES,
                    fila=num_fila,
                    campo="Principal",
                    valor=val_principal,
                )
            )

        if cod_prod and ubic_id:
            par = (cod_prod, ubic_id)
            if par in pares_asignacion_vistos:
                errores.append(
                    DetalleErrorInventario(
                        codigo=CodigoErrorInventario.ASIGNACION_DUPLICADA,
                        mensaje=(
                            f"Fila {num_fila} en '{HOJA_ASIGNACIONES}': asignación duplicada "
                            f"para el producto '{cod_prod}' y la ubicación '{ubic_id}'."
                        ),
                        hoja=HOJA_ASIGNACIONES,
                        fila=num_fila,
                        campo="Codigo_Producto,Ubicacion_ID",
                        valor=f"{cod_prod}:{ubic_id}",
                    )
                )
            else:
                pares_asignacion_vistos.add(par)

        if prod_valido:
            conteo_asignaciones_por_producto[cod_prod] = (
                conteo_asignaciones_por_producto.get(cod_prod, 0) + 1
            )
            if principal_normalizado == "Sí" and ubic_valida:
                conteo_principales_por_producto[cod_prod] = (
                    conteo_principales_por_producto.get(cod_prod, 0) + 1
                )
            elif principal_normalizado == "Sí" and not ubic_valida:
                # Aun cuando la ubicación sea inválida, contabilizamos la marca principal
                conteo_principales_por_producto[cod_prod] = (
                    conteo_principales_por_producto.get(cod_prod, 0) + 1
                )

    # 5. Reglas de cardinalidad por producto (al menos una ubicación y exactamente una principal)
    for cod_prod in orden_productos:
        total_asig = conteo_asignaciones_por_producto.get(cod_prod, 0)
        total_prin = conteo_principales_por_producto.get(cod_prod, 0)
        if total_asig == 0:
            errores.append(
                DetalleErrorInventario(
                    codigo=CodigoErrorInventario.PRODUCTO_SIN_ASIGNACION,
                    mensaje=(
                        f"El producto '{cod_prod}' no tiene ninguna ubicación asignada "
                        f"en '{HOJA_ASIGNACIONES}'."
                    ),
                    hoja=HOJA_ASIGNACIONES,
                    campo="Codigo_Producto",
                    valor=cod_prod,
                )
            )
        elif total_prin > 1:
            errores.append(
                DetalleErrorInventario(
                    codigo=CodigoErrorInventario.UBICACION_PRINCIPAL_MULTIPLE,
                    mensaje=(
                        f"El producto '{cod_prod}' tiene {total_prin} ubicaciones marcadas como "
                        f"principales; debe tener exactamente una."
                    ),
                    hoja=HOJA_ASIGNACIONES,
                    campo="Principal",
                    valor=cod_prod,
                )
            )
        elif total_prin == 0:
            errores.append(
                DetalleErrorInventario(
                    codigo=CodigoErrorInventario.SIN_UBICACION_PRINCIPAL,
                    mensaje=(
                        f"El producto '{cod_prod}' no tiene ninguna ubicación marcada como "
                        f"principal en '{HOJA_ASIGNACIONES}'."
                    ),
                    hoja=HOJA_ASIGNACIONES,
                    campo="Principal",
                    valor=cod_prod,
                )
            )

    if errores and lanzar_excepcion:
        _lanzar_excepcion_desde_errores(errores)

    return ResultadoValidacion(es_valido=len(errores) == 0, errores=errores)


def construir_inventario_datos(hojas: Mapping[str, pd.DataFrame]) -> InventarioDatos:
    """Valida las cuatro hojas, normaliza los textos y construye un objeto InventarioDatos."""
    validar_inventario(hojas, lanzar_excepcion=True)

    df_productos = _normalizar_dataframe_texto(
        hojas[HOJA_PRODUCTOS], ESQUEMA_HOJAS[HOJA_PRODUCTOS]
    )
    df_productos["Activo"] = df_productos["Activo"].map(normalizar_valor_activo)
    df_categorias = _normalizar_dataframe_texto(
        hojas[HOJA_CATEGORIAS], ESQUEMA_HOJAS[HOJA_CATEGORIAS]
    )
    df_ubicaciones = _normalizar_dataframe_texto(
        hojas[HOJA_UBICACIONES], ESQUEMA_HOJAS[HOJA_UBICACIONES]
    )
    df_ubicaciones["Activo"] = df_ubicaciones["Activo"].map(normalizar_valor_activo)
    df_asignaciones = _normalizar_dataframe_texto(
        hojas[HOJA_ASIGNACIONES], ESQUEMA_HOJAS[HOJA_ASIGNACIONES]
    )
    df_asignaciones["Principal"] = df_asignaciones["Principal"].map(normalizar_valor_principal)

    return InventarioDatos(
        productos=df_productos,
        categorias=df_categorias,
        ubicaciones=df_ubicaciones,
        asignaciones=df_asignaciones,
    )


def cargar_inventario_excel(origen: str | Path | bytes | BinaryIO) -> InventarioDatos:
    """Lee un archivo Excel desde ruta o flujo de bytes, preserva textos y valida su integridad.

    No modifica ni sobrescribe en ningún caso el archivo de origen.
    """
    buffer_lectura: BinaryIO
    nombre_origen = "<memoria>"

    if isinstance(origen, (str, Path)):
        ruta = Path(origen)
        nombre_origen = str(ruta)
        if not ruta.exists() or not ruta.is_file():
            raise ErrorArchivoInexistente(
                f"El archivo de inventario no existe en la ruta especificada: '{nombre_origen}'.",
                codigo=CodigoErrorInventario.ARCHIVO_INEXISTENTE,
            )
        try:
            contenido_bytes = ruta.read_bytes()
        except OSError as exc:
            raise ErrorArchivoCorrupto(
                f"No fue posible leer el archivo '{nombre_origen}': {exc}",
                codigo=CodigoErrorInventario.ARCHIVO_CORRUPTO,
            ) from exc
        if len(contenido_bytes) == 0:
            raise ErrorArchivoCorrupto(
                f"El archivo '{nombre_origen}' está vacío.",
                codigo=CodigoErrorInventario.ARCHIVO_CORRUPTO,
            )
        buffer_lectura = io.BytesIO(contenido_bytes)
    elif isinstance(origen, bytes):
        if len(origen) == 0:
            raise ErrorArchivoCorrupto(
                "El flujo de bytes proporcionado para el inventario está vacío.",
                codigo=CodigoErrorInventario.ARCHIVO_CORRUPTO,
            )
        buffer_lectura = io.BytesIO(origen)
    elif hasattr(origen, "read"):
        contenido = origen.read()
        if not isinstance(contenido, bytes) or len(contenido) == 0:
            raise ErrorArchivoCorrupto(
                "El objeto de lectura proporcionado está vacío o no contiene bytes válidos.",
                codigo=CodigoErrorInventario.ARCHIVO_CORRUPTO,
            )
        buffer_lectura = io.BytesIO(contenido)
    else:
        raise ErrorArchivoCorrupto(
            f"Tipo de origen no soportado para cargar inventario: {type(origen).__name__}.",
            codigo=CodigoErrorInventario.ARCHIVO_CORRUPTO,
        )

    try:
        hojas_leidas = pd.read_excel(
            buffer_lectura,
            sheet_name=None,
            dtype=str,
            keep_default_na=False,
            engine="openpyxl",
        )
    except (BadZipFile, InvalidFileException, ValueError, KeyError, OSError) as exc:
        raise ErrorArchivoCorrupto(
            f"El archivo '{nombre_origen}' está corrupto o no tiene un formato Excel (.xlsx) válido: {exc}",
            codigo=CodigoErrorInventario.ARCHIVO_CORRUPTO,
        ) from exc
    except Exception as exc:
        raise ErrorArchivoCorrupto(
            f"Error al procesar el archivo Excel '{nombre_origen}': {exc}",
            codigo=CodigoErrorInventario.ARCHIVO_CORRUPTO,
        ) from exc

    return construir_inventario_datos(hojas_leidas)


def formatear_ubicacion_fisica(ubicacion: Mapping[str, str]) -> str:
    """Devuelve una representación legible de la ubicación física."""
    return (
        f"Rack {ubicacion['Rack']} · Columna {ubicacion['Columna']} · "
        f"Nivel {ubicacion['Nivel']} · Posición {ubicacion['Posicion']}"
    )


def obtener_productos_estructurados(
    inventario: InventarioDatos, *, incluir_inactivos: bool = True
) -> list[dict[str, Any]]:
    """Prepara una lista estructurada de productos con categoría, ubicación principal y secundarias."""
    mapa_categorias: dict[str, str] = dict(
        zip(inventario.categorias["Categoria_ID"], inventario.categorias["Nombre_Categoria"])
    )

    mapa_ubicaciones: dict[str, dict[str, str]] = {}
    for _, fila_u in inventario.ubicaciones.iterrows():
        datos_u = {
            "ubicacion_id": fila_u["Ubicacion_ID"],
            "rack": fila_u["Rack"],
            "columna": fila_u["Columna"],
            "nivel": fila_u["Nivel"],
            "posicion": fila_u["Posicion"],
            "etiqueta": formatear_ubicacion_fisica(fila_u),
        }
        mapa_ubicaciones[fila_u["Ubicacion_ID"]] = datos_u

    asignaciones_por_producto: dict[str, list[tuple[str, str]]] = {}
    for _, fila_a in inventario.asignaciones.iterrows():
        cod = fila_a["Codigo_Producto"]
        asignaciones_por_producto.setdefault(cod, []).append(
            (fila_a["Ubicacion_ID"], fila_a["Principal"])
        )

    resultado: list[dict[str, Any]] = []
    for _, fila_p in inventario.productos.iterrows():
        cod_prod = fila_p["Codigo_Producto"]
        cat_id = fila_p["Categoria_ID"]
        val_activo_bruto = fila_p.get("Activo", "Sí") if "Activo" in fila_p else "Sí"
        try:
            activo_texto = normalizar_valor_activo(val_activo_bruto)
        except ErrorValorActivoInvalido:
            activo_texto = "Sí"
        es_activo = activo_texto == "Sí"

        if not incluir_inactivos and not es_activo:
            continue

        principal_info: dict[str, Any] | None = None
        secundarias_info: list[dict[str, Any]] = []

        for ubic_id, marca_principal in asignaciones_por_producto.get(cod_prod, []):
            base_ubic = dict(mapa_ubicaciones[ubic_id])
            es_principal = marca_principal == "Sí"
            base_ubic["es_principal"] = es_principal
            base_ubic["principal"] = marca_principal
            if es_principal:
                principal_info = base_ubic
            else:
                secundarias_info.append(base_ubic)

        todas = ([principal_info] if principal_info is not None else []) + secundarias_info
        resultado.append(
            {
                "codigo_producto": cod_prod,
                "nombre": fila_p["Nombre"],
                "categoria_id": cat_id,
                "nombre_categoria": mapa_categorias.get(cat_id, ""),
                "descripcion": fila_p["Descripcion"],
                "activo": activo_texto,
                "es_activo": es_activo,
                "ubicacion_principal": principal_info,
                "ubicaciones_secundarias": secundarias_info,
                "todas_las_ubicaciones": todas,
            }
        )

    return resultado


def obtener_productos_activos_estructurados(inventario: InventarioDatos) -> list[dict[str, Any]]:
    """Prepara la lista estructurada exclusivamente con los productos activos."""
    return obtener_productos_estructurados(inventario, incluir_inactivos=False)


def obtener_producto_por_codigo(
    inventario: InventarioDatos, codigo_producto: str, *, incluir_inactivos: bool = True
) -> dict[str, Any] | None:
    """Busca y devuelve la estructura completa de un producto por su código exacto."""
    codigo_limpio = _limpiar_texto(codigo_producto)
    for prod in obtener_productos_estructurados(
        inventario, incluir_inactivos=incluir_inactivos
    ):
        if prod["codigo_producto"] == codigo_limpio:
            return prod
    return None


def construir_tabla_consulta(inventario: InventarioDatos) -> pd.DataFrame:
    """Construye un DataFrame relacional unificado para facilitar consultas y auditorías."""
    df_prod_cat = inventario.productos.merge(
        inventario.categorias, on="Categoria_ID", how="left"
    )
    df_asig_ubic = inventario.asignaciones.merge(
        inventario.ubicaciones, on="Ubicacion_ID", how="left"
    )
    df_completo = df_prod_cat.merge(df_asig_ubic, on="Codigo_Producto", how="left")
    columnas_orden = [
        "Codigo_Producto",
        "Nombre",
        "Categoria_ID",
        "Nombre_Categoria",
        "Descripcion",
        "Ubicacion_ID",
        "Principal",
        "Rack",
        "Columna",
        "Nivel",
        "Posicion",
    ]
    return df_completo[columnas_orden].reset_index(drop=True)


def crear_inventario_demostrativo() -> InventarioDatos:
    """Crea en memoria un inventario demostrativo válido e identificado conforme a RF-031.

    Esta función opera exclusivamente en memoria y no escribe ningún archivo en disco.
    """
    df_categorias = pd.DataFrame(
        [
            {"Categoria_ID": "CAT01", "Nombre_Categoria": "Herramientas Manuales (Demostrativo)"},
            {"Categoria_ID": "CAT02", "Nombre_Categoria": "Tornillería y Fijación (Demostrativo)"},
            {"Categoria_ID": "CAT03", "Nombre_Categoria": "Electricidad e Iluminación (Demostrativo)"},
        ]
    )

    df_ubicaciones = pd.DataFrame(
        [
            {"Ubicacion_ID": "U001", "Rack": "R01", "Columna": "C01", "Nivel": "N02", "Posicion": "P03"},
            {"Ubicacion_ID": "U002", "Rack": "R02", "Columna": "C03", "Nivel": "N01", "Posicion": "P02"},
            {"Ubicacion_ID": "U003", "Rack": "R03", "Columna": "C02", "Nivel": "N03", "Posicion": "P01"},
            {"Ubicacion_ID": "U004", "Rack": "R01", "Columna": "C04", "Nivel": "N01", "Posicion": "P04"},
        ]
    )

    df_productos = pd.DataFrame(
        [
            {
                "Codigo_Producto": "MAR001",
                "Nombre": "Martillo de uña",
                "Categoria_ID": "CAT01",
                "Descripcion": f"{MARCA_DEMOSTRATIVA} Martillo para carpintería con mango ergonómico.",
            },
            {
                "Codigo_Producto": "TOR001",
                "Nombre": "Tornillo galvanizado",
                "Categoria_ID": "CAT02",
                "Descripcion": f"{MARCA_DEMOSTRATIVA} Tornillo de uso general resistente a la corrosión.",
            },
            {
                "Codigo_Producto": "000125",
                "Nombre": "Cinta aislante negra",
                "Categoria_ID": "CAT03",
                "Descripcion": f"{MARCA_DEMOSTRATIVA} Cinta aislante eléctrica de PVC de alta adherencia.",
            },
        ]
    )

    df_asignaciones = pd.DataFrame(
        [
            {"Codigo_Producto": "MAR001", "Ubicacion_ID": "U001", "Principal": "Sí"},
            {"Codigo_Producto": "MAR001", "Ubicacion_ID": "U002", "Principal": "No"},
            {"Codigo_Producto": "TOR001", "Ubicacion_ID": "U002", "Principal": "Sí"},
            {"Codigo_Producto": "000125", "Ubicacion_ID": "U003", "Principal": "Sí"},
            {"Codigo_Producto": "000125", "Ubicacion_ID": "U004", "Principal": "No"},
        ]
    )

    return construir_inventario_datos(
        {
            HOJA_PRODUCTOS: df_productos,
            HOJA_CATEGORIAS: df_categorias,
            HOJA_UBICACIONES: df_ubicaciones,
            HOJA_ASIGNACIONES: df_asignaciones,
        }
    )


def serializar_inventario_a_bytes(inventario: InventarioDatos) -> bytes:
    """Valida y serializa un objeto InventarioDatos a un flujo de bytes en formato .xlsx.

    Todas las celdas se almacenan explícitamente como texto ('@') para preservar ceros iniciales.
    """
    hojas = inventario.obtener_hojas()
    validar_inventario(hojas, lanzar_excepcion=True)

    wb = openpyxl.Workbook()
    # Eliminar la hoja por defecto creada por openpyxl
    hoja_inicial = wb.active
    if hoja_inicial is not None:
        wb.remove(hoja_inicial)

    for nombre_hoja in ORDEN_HOJAS:
        df = hojas[nombre_hoja]
        columnas = ESQUEMA_HOJAS[nombre_hoja]
        ws = wb.create_sheet(title=nombre_hoja)

        # Escribir encabezados como texto
        for col_idx, nombre_col in enumerate(columnas, start=1):
            celda_enc = ws.cell(row=1, column=col_idx, value=str(nombre_col))
            celda_enc.number_format = "@"
            celda_enc.data_type = "s"

        # Escribir filas como texto explícito
        for row_idx, (_, fila) in enumerate(df.iterrows(), start=2):
            for col_idx, nombre_col in enumerate(columnas, start=1):
                valor_str = _limpiar_texto(fila[nombre_col])
                celda = ws.cell(row=row_idx, column=col_idx, value=valor_str)
                celda.number_format = "@"
                celda.data_type = "s"

    salida = io.BytesIO()
    wb.save(salida)
    return salida.getvalue()


def guardar_inventario_en_ruta(
    inventario: InventarioDatos,
    ruta_destino: str | Path,
    permitir_sobrescritura: bool = False,
) -> Path:
    """Serializa y guarda el inventario en la ruta indicada protegiendo archivos existentes.

    Por defecto `permitir_sobrescritura=False`, impidiendo sobrescribir un archivo ya existente.
    """
    ruta = Path(ruta_destino)
    if ruta.exists() and not permitir_sobrescritura:
        raise ErrorProteccionInventario(
            f"Protección activa: el archivo '{ruta}' ya existe y no puede sobrescribirse automáticamente.",
            codigo=CodigoErrorInventario.PROTECCION_INVENTARIO_EXISTENTE,
        )

    contenido = serializar_inventario_a_bytes(inventario)
    ruta.parent.mkdir(parents=True, exist_ok=True)
    ruta.write_bytes(contenido)
    return ruta


def preparar_inicializacion_segura(
    estado_remoto: EstadoExistenciaRemota | str,
    ruta_local_control: str | Path | None = None,
) -> bytes:
    """Prepara en memoria los bytes del inventario demostrativo solo cuando se confirma su inexistencia.

    Reglas de seguridad (RF-029, RF-030, CA-020, CA-021):
    - Si `estado_remoto` es `EXISTE`, rechaza la inicialización para no sobrescribir el maestro.
    - Si `estado_remoto` es `DESCONOCIDO_ERROR` o cualquier valor distinto de `NO_EXISTE`,
      detiene la inicialización: un fallo de conexión o incertidumbre nunca equivale a inexistencia.
    - Si se proporciona `ruta_local_control` y dicho archivo ya existe localmente, también rechaza
      la operación para evitar reemplazos accidentales.
    - Nunca escribe `03_Aplicacion/inventario.xlsx` de forma automática.
    """
    try:
        estado = (
            estado_remoto
            if isinstance(estado_remoto, EstadoExistenciaRemota)
            else EstadoExistenciaRemota(str(estado_remoto))
        )
    except ValueError as exc:
        raise ErrorProteccionInventario(
            f"Estado remoto no reconocido ('{estado_remoto}'). "
            "No es posible confirmar la inexistencia del inventario maestro.",
            codigo=CodigoErrorInventario.ESTADO_REMOTO_DESCONOCIDO,
        ) from exc

    if estado == EstadoExistenciaRemota.EXISTE:
        raise ErrorProteccionInventario(
            "El archivo maestro 'inventario.xlsx' ya existe en GitHub. "
            "Se prohíbe recrearlo o sobrescribirlo automáticamente.",
            codigo=CodigoErrorInventario.PROTECCION_INVENTARIO_EXISTENTE,
        )

    if estado != EstadoExistenciaRemota.NO_EXISTE:
        raise ErrorProteccionInventario(
            "No se pudo determinar con seguridad si el archivo maestro existe en GitHub "
            "(error de conexión, permisos o estado desconocido). Inicialización detenida.",
            codigo=CodigoErrorInventario.ESTADO_REMOTO_DESCONOCIDO,
        )

    if ruta_local_control is not None and Path(ruta_local_control).exists():
        raise ErrorProteccionInventario(
            f"El archivo local '{ruta_local_control}' ya existe. "
            "No se sobrescribirá automáticamente.",
            codigo=CodigoErrorInventario.PROTECCION_INVENTARIO_EXISTENTE,
        )

    inventario_demo = crear_inventario_demostrativo()
    return serializar_inventario_a_bytes(inventario_demo)

