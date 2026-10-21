"""Pruebas automatizadas para la persistencia administrativa del inventario Excel en GitHub.

Proyecto: FerreteríaMS v0
Tarea: FERMS-V0-P10
Módulos probados: `03_Aplicacion/persistencia_github.py`, `03_Aplicacion/administracion_productos.py`
y `03_Aplicacion/app.py`.
Requisitos cubiertos: RF-012, RF-013, RF-014, RF-015, RF-016, RF-017, RF-018, RF-019,
RF-020, RF-021, RF-022, RF-023, RF-024, RF-025, RF-026, RF-027, RF-028, RF-029,
RF-030, RF-031, RF-035, RF-036, RF-037, RF-038 y Sección 6.3.
Criterios cubiertos: CA-012, CA-013, CA-014, CA-015, CA-016, CA-017, CA-018, CA-019,
CA-020, CA-021, CA-022, CA-023, CA-024, CA-025, CA-026, CA-042, CA-044, CA-045,
CA-046, CA-047, CA-048, CA-049, CA-050, CA-051, CA-054, CA-055, CA-056, CA-057 y CA-058.

Todas las pruebas de esta suite se ejecutan con mocks y simulaciones deterministas en memoria:
- No realizan llamadas HTTP PUT, POST, PATCH, DELETE ni GET reales hacia GitHub.
- No utilizan credenciales ni tokens reales.
- No crean ni modifican `03_Aplicacion/inventario.xlsx` localmente ni en GitHub.
"""

from __future__ import annotations

import base64
import hashlib
import inspect
import io
import json
import os
import re
import secrets
import socket
import subprocess
import sys
import unittest
import urllib.error
import urllib.request
from pathlib import Path
from typing import Any, Mapping
from unittest.mock import patch

import openpyxl
import pandas as pd
from streamlit.testing.v1 import AppTest

RAIZ_PROYECTO = Path(__file__).resolve().parent.parent
RUTA_APLICACION = RAIZ_PROYECTO / "03_Aplicacion"
RUTA_APP = RUTA_APLICACION / "app.py"
RUTA_ADMIN_PRODUCTOS = RUTA_APLICACION / "administracion_productos.py"
RUTA_AUTENTICACION = RUTA_APLICACION / "autenticacion.py"
RUTA_PERSISTENCIA = RUTA_APLICACION / "persistencia_github.py"
RUTA_INVENTARIO_XLSX = RUTA_APLICACION / "inventario.xlsx"
BYTES_INVENTARIO_XLSX_INICIAL = (
    RUTA_INVENTARIO_XLSX.read_bytes() if RUTA_INVENTARIO_XLSX.exists() else None
)

if str(RUTA_APLICACION) not in sys.path:
    sys.path.insert(0, str(RUTA_APLICACION))

import administracion_productos as admin_prod  # noqa: E402
import app as app_ferms  # noqa: E402
import autenticacion  # noqa: E402
import busqueda  # noqa: E402
import github_inventario as gh  # noqa: E402
import inventario as motor_inv  # noqa: E402
import persistencia_github as persist_gh  # noqa: E402
from administracion_productos import (  # noqa: E402
    CLAVE_ESPACIO_TRABAJO_ADMIN,
    CLAVE_FORMULARIO_BORRADOR_ADMIN,
    AsignacionUbicacionEntrada,
    CodigoErrorAdministrativo,
    ErrorAdministracionInventario,
    ErrorValidacionAdministrativa,
    EspacioTrabajoAdminInventario,
    EstadoInventarioAdmin,
    cambiar_ubicacion_principal_local,
    cargar_inventario_administrativo_en_sesion,
    clonar_inventario_datos,
    consultar_producto_administrativo_por_codigo,
    consultar_productos_administrativos,
    crear_producto_local,
    editar_producto_local,
    eliminar_asignacion_ubicacion_local,
    guardar_borrador_formulario_admin,
    inicializar_espacio_trabajo_desde_inventario,
    inventarios_son_identicos,
    restaurar_borrador_en_espacio_trabajo_admin,
)
from autenticacion import (  # noqa: E402
    TIEMPO_EXPIRACION_INACTIVIDAD_SEGUNDOS,
    CodigoResultadoAutenticacion,
    ContextoDispositivo,
    ErrorAutorizacionAdmin,
    configurar_reloj_servidor_pruebas,
    es_hash_contrasena_valido,
    generar_hash_contrasena,
    iniciar_sesion_admin,
    obtener_borrador_admin,
    reiniciar_estado_seguridad_servidor,
    tiene_borrador_pendiente_en_sesion,
)
from github_inventario import (  # noqa: E402
    CONFIGURACION_POR_DEFECTO,
    CategoriaResultadoConsulta,
    CodigoErrorGitHub,
    ConfiguracionRepositorioGitHub,
    RespuestaHTTPGitHub,
    ResultadoExistenciaRemota,
    ResultadoLecturaInventarioRemoto,
)
from inventario import (  # noqa: E402
    ESQUEMA_HOJAS as ESQUEMA_COLUMNAS,
    HOJA_ASIGNACIONES,
    HOJA_CATEGORIAS,
    HOJA_PRODUCTOS,
    HOJA_UBICACIONES,
    ORDEN_HOJAS,
    EstadoExistenciaRemota,
    InventarioDatos,
    cargar_inventario_excel,
    construir_inventario_datos,
    crear_inventario_demostrativo,
    serializar_inventario_a_bytes,
)
from persistencia_github import (  # noqa: E402
    CLAVE_CONFIRMACION_GUARDADO_GITHUB,
    CLAVE_CONFLICTO_SHA_ADMIN,
    CLAVE_HUELLA_ULTIMA_PUBLICACION_ADMIN,
    CLAVE_ULTIMA_PUBLICACION_GITHUB_ADMIN,
    ESCRITURA_REAL_GITHUB_HABILITADA,
    MENSAJE_ARCHIVO_REMOTO_YA_EXISTE,
    MENSAJE_CONFIRMACION_EXPLICITA_REQUERIDA,
    MENSAJE_CONFLICTO_SHA_DETECTADO,
    MENSAJE_PUBLICACION_ACTUALIZACION_EXITOSA,
    MENSAJE_PUBLICACION_DUPLICADA_BLOQUEADA,
    MENSAJE_PUBLICACION_INICIALIZACION_EXITOSA,
    CodigoErrorPersistenciaGitHub,
    ErrorConfirmacionPublicacionRequerida,
    ErrorConflictoVersionShaGitHub,
    ErrorCredencialesEscrituraGitHub,
    ErrorDestinoEscrituraNoAutorizadoGitHub,
    ErrorEscrituraRealNoAutorizadaGitHub,
    ErrorInventarioRemotoYaExiste,
    ErrorLimiteSolicitudesEscrituraGitHub,
    ErrorPersistenciaGitHub,
    ErrorPublicacionDatosDemostrativosNoAutorizada,
    ErrorPublicacionDuplicadaGitHub,
    ErrorRedOTimeoutEscrituraGitHub,
    ErrorRepositorioORamaEscrituraGitHub,
    ErrorRespuestaInesperadaEscrituraGitHub,
    ErrorSerializacionInventarioPersistente,
    ErrorServidorEscrituraGitHub,
    ErrorShaBaseFaltanteGitHub,
    ErrorSinCambiosPendientesPublicacion,
    ErrorTokenEscrituraFaltanteGitHub,
    EstadoConflictoShaAdmin,
    EstadoResultadoPublicacionGitHub,
    ModoPublicacionGitHub,
    construir_solicitud_put_contents_github,
    ejecutar_solicitud_put_github,
    interpretar_respuesta_put_github,
    publicar_cambios_administrativos_github,
    recargar_inventario_remoto_tras_conflicto,
    serializar_y_verificar_inventario_xlsx,
    validar_destino_escritura_github,
    verificar_inventario_autorizado_para_publicacion,
)

SHA_BASE_SIMULADO = "a1b2c3d4e5f60718293a4b5c6d7e8f9012345678"
SHA_NUEVO_SIMULADO = "f9e8d7c6b5a4938271605f4e3d2c1b0a98765432"
SHA_COMMIT_SIMULADO = "c0ffee1234567890abcdef1234567890abcdef12"
SHA_REMOTO_CONCURRENTE = "deadbeef99998888777766665555444433332222"
TOKEN_ESCRITURA_SIMULADO = "ghp_TokenSimuladoEscrituraP10NoReal12345678"


class RelojSimuladoP10:
    """Reloj determinista para simular pasos de tiempo sin esperas reales."""

    def __init__(self, epoch_inicial: float = 1_850_000_000.0) -> None:
        self.actual = float(epoch_inicial)

    def __call__(self) -> float:
        return self.actual

    def avanzar(self, segundos: float) -> float:
        self.actual += float(segundos)
        return self.actual


def _construir_inventario_base_p10() -> InventarioDatos:
    """Construye un inventario válido en memoria con códigos que incluyen ceros iniciales."""
    df_categorias = pd.DataFrame(
        [
            {"Categoria_ID": "CAT01", "Nombre_Categoria": "Herramientas Manuales"},
            {"Categoria_ID": "01", "Nombre_Categoria": "Fijaciones y Tornillería"},
            {"Categoria_ID": "CAT03", "Nombre_Categoria": "Medición y Nivelación"},
        ]
    )
    df_ubicaciones = pd.DataFrame(
        [
            {
                "Ubicacion_ID": "U001",
                "Rack": "A",
                "Columna": "1",
                "Nivel": "1",
                "Posicion": "01",
            },
            {
                "Ubicacion_ID": "01",
                "Rack": "A",
                "Columna": "2",
                "Nivel": "1",
                "Posicion": "02",
            },
            {
                "Ubicacion_ID": "A-01",
                "Rack": "B",
                "Columna": "3",
                "Nivel": "2",
                "Posicion": "04",
            },
            {
                "Ubicacion_ID": "U004",
                "Rack": "C",
                "Columna": "1",
                "Nivel": "3",
                "Posicion": "05",
            },
        ]
    )
    df_productos = pd.DataFrame(
        [
            {
                "Codigo_Producto": "0001",
                "Nombre": "Tornillo Drywall 6x1",
                "Categoria_ID": "01",
                "Descripcion": "Tornillo fosfatado punta fina.",
            },
            {
                "Codigo_Producto": "01",
                "Nombre": "Arandela Plana Zincada 1/4",
                "Categoria_ID": "01",
                "Descripcion": "Arandela metálica anticorrosiva.",
            },
            {
                "Codigo_Producto": "A-01",
                "Nombre": "Martillo de Uña 16 oz",
                "Categoria_ID": "CAT01",
                "Descripcion": "Martillo con mango de fibra de vidrio.",
            },
        ]
    )
    df_asignaciones = pd.DataFrame(
        [
            {"Codigo_Producto": "0001", "Ubicacion_ID": "U001", "Principal": "Sí"},
            {"Codigo_Producto": "0001", "Ubicacion_ID": "01", "Principal": "No"},
            {"Codigo_Producto": "01", "Ubicacion_ID": "01", "Principal": "Sí"},
            {"Codigo_Producto": "A-01", "Ubicacion_ID": "A-01", "Principal": "Sí"},
            {"Codigo_Producto": "A-01", "Ubicacion_ID": "U004", "Principal": "No"},
        ]
    )
    return construir_inventario_datos(
        {
            HOJA_PRODUCTOS: df_productos,
            HOJA_CATEGORIAS: df_categorias,
            HOJA_UBICACIONES: df_ubicaciones,
            HOJA_ASIGNACIONES: df_asignaciones,
        }
    )


def _respuesta_http_json(
    status_code: int,
    datos: Any,
    encabezados: Mapping[str, str] | None = None,
    url: str = "https://api.github.com/repos/JunyMSAI/FerreteriaMS/contents/03_Aplicacion/inventario.xlsx",
) -> RespuestaHTTPGitHub:
    """Construye una `RespuestaHTTPGitHub` simulada con cuerpo JSON."""
    cuerpo = json.dumps(datos, ensure_ascii=False).encode("utf-8")
    return RespuestaHTTPGitHub(
        status_code=status_code,
        cuerpo=cuerpo,
        encabezados=dict(encabezados or {"Content-Type": "application/json"}),
        url=url,
    )


def _respuesta_put_exitosa_simulada(
    *,
    nuevo_sha_archivo: str = SHA_NUEVO_SIMULADO,
    sha_commit: str = SHA_COMMIT_SIMULADO,
    status_code: int = 200,
    ruta: str = "03_Aplicacion/inventario.xlsx",
) -> RespuestaHTTPGitHub:
    """Construye una respuesta HTTP 200/201 simulada de GitHub Contents API tras un PUT exitoso."""
    return _respuesta_http_json(
        status_code,
        {
            "content": {
                "name": "inventario.xlsx",
                "path": ruta,
                "sha": nuevo_sha_archivo,
                "size": 5120,
                "type": "file",
            },
            "commit": {
                "sha": sha_commit,
                "message": "Actualización administrativa de inventario.xlsx (FerreteríaMS v0)",
            },
        },
    )


def _existencia_remota_existe(
    sha_github: str = SHA_BASE_SIMULADO,
) -> ResultadoExistenciaRemota:
    """Simula una consulta previa donde `inventario.xlsx` existe con el SHA indicado."""
    return ResultadoExistenciaRemota(
        estado=EstadoExistenciaRemota.EXISTE,
        propietario="JunyMSAI",
        repositorio="FerreteriaMS",
        rama="main",
        ruta_archivo="03_Aplicacion/inventario.xlsx",
        sha_github=sha_github,
        tamano_bytes=4096,
        categoria_consulta=CategoriaResultadoConsulta.INVENTARIO_DISPONIBLE,
        mensaje="El archivo maestro existe en GitHub.",
    )


def _existencia_remota_no_existe_confirmada() -> ResultadoExistenciaRemota:
    """Simula una consulta previa donde GitHub confirma inequívocamente que el archivo no existe."""
    return ResultadoExistenciaRemota(
        estado=EstadoExistenciaRemota.NO_EXISTE,
        propietario="JunyMSAI",
        repositorio="FerreteriaMS",
        rama="main",
        ruta_archivo="03_Aplicacion/inventario.xlsx",
        sha_github=None,
        tamano_bytes=None,
        codigo_error=CodigoErrorGitHub.ARCHIVO_REMOTO_INEXISTENTE,
        categoria_consulta=CategoriaResultadoConsulta.INVENTARIO_NO_INICIALIZADO,
        mensaje="GitHub confirmó que el archivo maestro no existe en la rama main.",
    )


class RepositorioRemotoSimuladoMemoria:
    """Simulador en memoria del estado del archivo maestro en GitHub para flujos completos."""

    def __init__(
        self,
        inventario_inicial: InventarioDatos | None,
        sha_inicial: str | None = SHA_BASE_SIMULADO,
    ) -> None:
        self.inventario_remoto: InventarioDatos | None = (
            clonar_inventario_datos(inventario_inicial)
            if inventario_inicial is not None
            else None
        )
        self.bytes_remotos: bytes | None = (
            serializar_inventario_a_bytes(self.inventario_remoto)
            if self.inventario_remoto is not None
            else None
        )
        self.sha_actual: str | None = (
            sha_inicial if self.inventario_remoto is not None else None
        )
        self.contador_commits: int = 0
        self.historial_puts: list[dict[str, Any]] = []

    def consultar_existencia(self, **kwargs: Any) -> ResultadoExistenciaRemota:
        if self.inventario_remoto is None or self.sha_actual is None:
            return _existencia_remota_no_existe_confirmada()
        return _existencia_remota_existe(sha_github=self.sha_actual)

    def descargar_remoto(self, **kwargs: Any) -> ResultadoLecturaInventarioRemoto:
        if (
            self.inventario_remoto is None
            or self.bytes_remotos is None
            or self.sha_actual is None
        ):
            raise gh.ErrorArchivoRemotoInexistente("No existe inventario.xlsx en remoto.")
        return ResultadoLecturaInventarioRemoto(
            propietario="JunyMSAI",
            repositorio="FerreteriaMS",
            rama="main",
            ruta_archivo="03_Aplicacion/inventario.xlsx",
            sha_github=self.sha_actual,
            sha256_local=hashlib.sha256(self.bytes_remotos).hexdigest(),
            tamano_bytes=len(self.bytes_remotos),
            inventario=clonar_inventario_datos(self.inventario_remoto),
            categoria_consulta=CategoriaResultadoConsulta.INVENTARIO_DISPONIBLE,
        )

    def ejecutar_put(
        self,
        url: str,
        encabezados: Mapping[str, str],
        cuerpo_json_bytes: bytes,
        timeout: float = 15.0,
        **kwargs: Any,
    ) -> RespuestaHTTPGitHub:
        payload = json.loads(cuerpo_json_bytes.decode("utf-8"))
        self.historial_puts.append(
            {
                "url": url,
                "encabezados": dict(encabezados),
                "payload": payload,
                "timeout": timeout,
            }
        )
        sha_enviado = payload.get("sha")
        if self.sha_actual is not None:
            if sha_enviado != self.sha_actual:
                return _respuesta_http_json(
                    409,
                    {
                        "message": (
                            f"03_Aplicacion/inventario.xlsx does not match {self.sha_actual}"
                        )
                    },
                    url=url,
                )
        else:
            if sha_enviado is not None:
                return _respuesta_http_json(
                    422,
                    {"message": "Invalid request: sha not expected for new file"},
                    url=url,
                )

        bytes_nuevos = base64.b64decode(payload["content"], validate=True)
        inv_nuevo = cargar_inventario_excel(bytes_nuevos)
        self.contador_commits += 1
        es_creacion = self.sha_actual is None
        nuevo_sha = hashlib.sha1(
            f"blob:{self.contador_commits}:{len(bytes_nuevos)}".encode("utf-8")
        ).hexdigest()
        commit_sha = hashlib.sha1(
            f"commit:{self.contador_commits}:{payload.get('message', '')}".encode("utf-8")
        ).hexdigest()
        self.bytes_remotos = bytes_nuevos
        self.inventario_remoto = inv_nuevo
        self.sha_actual = nuevo_sha
        return _respuesta_put_exitosa_simulada(
            nuevo_sha_archivo=nuevo_sha,
            sha_commit=commit_sha,
            status_code=201 if es_creacion else 200,
        )


class TestPrompt10PersistenciaInventarioGitHub(unittest.TestCase):
    """Suite de 40 escenarios obligatorios del Prompt 10 (Persistencia en GitHub Contents API)."""

    USUARIO_PRUEBA: str = ""
    CONTRASENA_PRUEBA: str = ""
    HASH_PRUEBA_ARGON2ID: str = ""
    SECRETOS_VALIDOS: dict[str, str] = {}

    @classmethod
    def setUpClass(cls) -> None:
        cls.USUARIO_PRUEBA = f"admin_p10_{secrets.token_hex(4)}"
        cls.CONTRASENA_PRUEBA = f"ClaveSeguraP10_{secrets.token_urlsafe(10)}"
        cls.HASH_PRUEBA_ARGON2ID = generar_hash_contrasena(cls.CONTRASENA_PRUEBA)
        cls.SECRETOS_VALIDOS = {
            "FERRETERIAMS_ADMIN_USER": cls.USUARIO_PRUEBA,
            "FERRETERIAMS_ADMIN_PASSWORD_HASH": cls.HASH_PRUEBA_ARGON2ID,
            "GITHUB_TOKEN": TOKEN_ESCRITURA_SIMULADO,
        }
        cls.inventario_base = _construir_inventario_base_p10()

    def setUp(self) -> None:
        reiniciar_estado_seguridad_servidor()
        self.reloj = RelojSimuladoP10()
        configurar_reloj_servidor_pruebas(self.reloj)
        self.secretos_validos = dict(self.SECRETOS_VALIDOS)

        def _verificar_rapido(contrasena_plana: Any, hash_almacenado: str) -> bool:
            if not isinstance(contrasena_plana, str) or not contrasena_plana:
                return False
            if not es_hash_contrasena_valido(hash_almacenado):
                return False
            return (
                contrasena_plana == self.CONTRASENA_PRUEBA
                and hash_almacenado == self.HASH_PRUEBA_ARGON2ID
            )

        self.patcher_argon2 = patch.object(
            autenticacion,
            "verificar_contrasena_hash",
            side_effect=_verificar_rapido,
        )
        self.patcher_argon2.start()

    def tearDown(self) -> None:
        self.patcher_argon2.stop()
        reiniciar_estado_seguridad_servidor()

    def _crear_sesion_y_espacio_admin(
        self,
        *,
        inventario: InventarioDatos | None = None,
        sha_base: str | None = SHA_BASE_SIMULADO,
        origen_remoto: bool = True,
        es_copia_demostrativa: bool | None = None,
        inicializacion_oficial_autorizada: bool = False,
    ) -> tuple[dict[str, Any], EspacioTrabajoAdminInventario]:
        estado: dict[str, Any] = {}
        app_ferms.inicializar_estado_sesion(estado)
        res_login = iniciar_sesion_admin(
            usuario=self.USUARIO_PRUEBA,
            contrasena=self.CONTRASENA_PRUEBA,
            estado=estado,
            fuente_secretos=self.secretos_validos,
            contexto_dispositivo=ContextoDispositivo.ESCRITORIO,
            reloj=self.reloj,
        )
        self.assertTrue(res_login.autenticado)
        espacio = inicializar_espacio_trabajo_desde_inventario(
            estado,
            inventario if inventario is not None else self.inventario_base,
            sha_base_remoto=sha_base,
            origen_remoto_confirmado=origen_remoto,
            es_copia_demostrativa=es_copia_demostrativa,
            inicializacion_oficial_autorizada=inicializacion_oficial_autorizada,
            contexto_dispositivo=ContextoDispositivo.ESCRITORIO,
            reloj=self.reloj,
        )
        return estado, espacio

    # ------------------------------------------------------------------
    # 1. Serialización del inventario a bytes .xlsx en memoria
    # ------------------------------------------------------------------
    def test_01_serializacion_inventario_a_bytes_xlsx_en_memoria(self) -> None:
        paquete = serializar_y_verificar_inventario_xlsx(self.inventario_base)
        self.assertIsInstance(paquete.bytes_xlsx, bytes)
        self.assertGreater(paquete.tamano_bytes, 0)
        self.assertEqual(paquete.tamano_bytes, len(paquete.bytes_xlsx))
        self.assertEqual(
            paquete.sha256_bytes,
            hashlib.sha256(paquete.bytes_xlsx).hexdigest(),
        )
        # Firma PK\x03\x04 de archivo .xlsx (contenedor ZIP OpenXML)
        self.assertTrue(paquete.bytes_xlsx.startswith(b"PK\x03\x04"))
        self.assertTrue(RUTA_INVENTARIO_XLSX.is_file())
        self.assertEqual(RUTA_INVENTARIO_XLSX.read_bytes(), BYTES_INVENTARIO_XLSX_INICIAL)

    # ------------------------------------------------------------------
    # 2. Preservación exacta de las cuatro hojas (Productos, Categorías, Ubicaciones, Asignaciones)
    # ------------------------------------------------------------------
    def test_02_preservacion_exacta_cuatro_hojas(self) -> None:
        paquete = serializar_y_verificar_inventario_xlsx(self.inventario_base)
        wb = openpyxl.load_workbook(io.BytesIO(paquete.bytes_xlsx), read_only=True)
        try:
            self.assertEqual(
                tuple(wb.sheetnames),
                ("Productos", "Categorías", "Ubicaciones", "Asignaciones"),
            )
            self.assertEqual(tuple(wb.sheetnames), ORDEN_HOJAS)
        finally:
            wb.close()

    # ------------------------------------------------------------------
    # 3. Preservación del orden de columnas en cada hoja
    # ------------------------------------------------------------------
    def test_03_preservacion_orden_columnas_en_cada_hoja(self) -> None:
        paquete = serializar_y_verificar_inventario_xlsx(self.inventario_base)
        wb = openpyxl.load_workbook(io.BytesIO(paquete.bytes_xlsx), data_only=True)
        try:
            for nombre_hoja, columnas_esperadas in ESQUEMA_COLUMNAS.items():
                ws = wb[nombre_hoja]
                encabezado = tuple(cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1)))
                self.assertEqual(encabezado, tuple(columnas_esperadas))
        finally:
            wb.close()

    # ------------------------------------------------------------------
    # 4. Preservación de códigos como texto con ceros iniciales (0001, 01, A-01)
    # ------------------------------------------------------------------
    def test_04_preservacion_codigos_como_texto_con_ceros_iniciales(self) -> None:
        paquete = serializar_y_verificar_inventario_xlsx(self.inventario_base)
        inv_releido = paquete.inventario_verificado
        codigos_prod = list(inv_releido.productos["Codigo_Producto"])
        self.assertIn("0001", codigos_prod)
        self.assertIn("01", codigos_prod)
        self.assertIn("A-01", codigos_prod)
        self.assertNotIn("1", codigos_prod)

        codigos_cat = list(inv_releido.categorias["Categoria_ID"])
        self.assertIn("01", codigos_cat)

        codigos_ubic = list(inv_releido.ubicaciones["Ubicacion_ID"])
        self.assertIn("01", codigos_ubic)
        self.assertIn("A-01", codigos_ubic)

        posiciones = list(inv_releido.ubicaciones["Posicion"])
        self.assertIn("01", posiciones)
        self.assertIn("02", posiciones)

        # Verificar además formato '@' y tipo texto en celdas de openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(paquete.bytes_xlsx), data_only=False)
        try:
            celda_cod = wb[HOJA_PRODUCTOS].cell(row=2, column=1)
            self.assertEqual(celda_cod.value, "0001")
            self.assertEqual(celda_cod.data_type, "s")
            self.assertEqual(celda_cod.number_format, "@")
        finally:
            wb.close()

    # ------------------------------------------------------------------
    # 5. Verificación round-trip leyendo los bytes generados con cargar_inventario_excel()
    # ------------------------------------------------------------------
    def test_05_verificacion_roundtrip_con_cargar_inventario_excel(self) -> None:
        paquete = serializar_y_verificar_inventario_xlsx(self.inventario_base)
        inv_desde_bytes = cargar_inventario_excel(paquete.bytes_xlsx)
        self.assertTrue(inventarios_son_identicos(self.inventario_base, inv_desde_bytes))
        self.assertTrue(
            inventarios_son_identicos(paquete.inventario_verificado, inv_desde_bytes)
        )

    # ------------------------------------------------------------------
    # 6. Rechazo de serialización cuando el inventario contiene errores de esquema o integridad
    # ------------------------------------------------------------------
    def test_06_rechazo_serializacion_con_errores_esquema_o_integridad(self) -> None:
        # Caso A: Objeto que no es InventarioDatos
        with self.assertRaises(ErrorSerializacionInventarioPersistente) as ctx_tipo:
            serializar_y_verificar_inventario_xlsx({"invalido": True})  # type: ignore[arg-type]
        self.assertEqual(
            ctx_tipo.exception.codigo,
            CodigoErrorPersistenciaGitHub.SERIALIZACION_INVALIDA,
        )

        # Caso B: InventarioDatos corrompido en memoria (producto sin ubicación principal)
        inv_corrupto = clonar_inventario_datos(self.inventario_base)
        inv_corrupto.asignaciones["Principal"] = "No"
        with self.assertRaises(ErrorSerializacionInventarioPersistente) as ctx_esquema:
            serializar_y_verificar_inventario_xlsx(inv_corrupto)
        self.assertEqual(
            ctx_esquema.exception.codigo,
            CodigoErrorPersistenciaGitHub.SERIALIZACION_INVALIDA,
        )

        # Caso C: Fallo detectado en la verificación round-trip (diferencia entre entrada y relectura)
        inv_diferente = crear_inventario_demostrativo()
        with self.assertRaises(ErrorSerializacionInventarioPersistente) as ctx_rt:
            serializar_y_verificar_inventario_xlsx(
                self.inventario_base,
                lector_roundtrip=lambda _b: inv_diferente,
            )
        self.assertEqual(
            ctx_rt.exception.codigo,
            CodigoErrorPersistenciaGitHub.VERIFICACION_ROUNDTRIP_FALLIDA,
        )

    # ------------------------------------------------------------------
    # 7. Construcción correcta del payload JSON para GitHub Contents API
    # ------------------------------------------------------------------
    def test_07_construccion_correcta_payload_json_github_contents_api(self) -> None:
        solicitud, _ = construir_solicitud_put_contents_github(
            self.inventario_base,
            modo_publicacion=ModoPublicacionGitHub.ACTUALIZACION_ARCHIVO_EXISTENTE,
            sha_base_remoto=SHA_BASE_SIMULADO,
            mensaje_commit="Actualización controlada de prueba P10",
            token=TOKEN_ESCRITURA_SIMULADO,
        )
        self.assertEqual(solicitud.metodo, "PUT")
        payload_decodificado = json.loads(solicitud.cuerpo_json_bytes.decode("utf-8"))
        self.assertEqual(
            set(payload_decodificado.keys()),
            {"message", "content", "branch", "sha"},
        )
        self.assertEqual(
            payload_decodificado["message"],
            "Actualización controlada de prueba P10",
        )
        self.assertEqual(payload_decodificado["branch"], "main")
        self.assertEqual(payload_decodificado["sha"], SHA_BASE_SIMULADO)

    # ------------------------------------------------------------------
    # 8. Codificación Base64 válida del archivo .xlsx
    # ------------------------------------------------------------------
    def test_08_codificacion_base64_valida_del_archivo_xlsx(self) -> None:
        solicitud, paquete = construir_solicitud_put_contents_github(
            self.inventario_base,
            modo_publicacion=ModoPublicacionGitHub.ACTUALIZACION_ARCHIVO_EXISTENTE,
            sha_base_remoto=SHA_BASE_SIMULADO,
            token=TOKEN_ESCRITURA_SIMULADO,
        )
        b64_str = solicitud.payload["content"]
        self.assertIsInstance(b64_str, str)
        bytes_decodificados = base64.b64decode(b64_str, validate=True)
        self.assertEqual(bytes_decodificados, paquete.bytes_xlsx)
        inv_reconstruido = cargar_inventario_excel(bytes_decodificados)
        self.assertTrue(inventarios_son_identicos(self.inventario_base, inv_reconstruido))

    # ------------------------------------------------------------------
    # 9. Inclusión de branch="main" y ruta 03_Aplicacion/inventario.xlsx
    # ------------------------------------------------------------------
    def test_09_inclusion_branch_main_y_ruta_oficial_y_rechazo_destinos_no_autorizados(
        self,
    ) -> None:
        solicitud, _ = construir_solicitud_put_contents_github(
            self.inventario_base,
            modo_publicacion=ModoPublicacionGitHub.ACTUALIZACION_ARCHIVO_EXISTENTE,
            sha_base_remoto=SHA_BASE_SIMULADO,
            token=TOKEN_ESCRITURA_SIMULADO,
        )
        self.assertEqual(
            solicitud.url,
            "https://api.github.com/repos/JunyMSAI/FerreteriaMS/contents/03_Aplicacion/inventario.xlsx",
        )
        self.assertEqual(solicitud.payload["branch"], "main")

        # Rechazar propietario, repositorio, rama, ruta o protocolo distintos de los oficiales
        configs_invalidas = [
            ConfiguracionRepositorioGitHub(propietario="OtroOwner"),
            ConfiguracionRepositorioGitHub(repositorio="OtroRepo"),
            ConfiguracionRepositorioGitHub(rama="develop"),
            ConfiguracionRepositorioGitHub(ruta_archivo="otro_archivo.xlsx"),
            ConfiguracionRepositorioGitHub(url_base_api="http://api.github.com"),
        ]
        for cfg_invalida in configs_invalidas:
            with self.assertRaises(ErrorDestinoEscrituraNoAutorizadoGitHub):
                validar_destino_escritura_github(cfg_invalida)

    # ------------------------------------------------------------------
    # 10. Inclusión del sha vigente al actualizar un archivo existente
    # ------------------------------------------------------------------
    def test_10_inclusion_sha_vigente_al_actualizar_archivo_existente(self) -> None:
        solicitud, _ = construir_solicitud_put_contents_github(
            self.inventario_base,
            modo_publicacion=ModoPublicacionGitHub.ACTUALIZACION_ARCHIVO_EXISTENTE,
            sha_base_remoto=SHA_BASE_SIMULADO,
            token=TOKEN_ESCRITURA_SIMULADO,
        )
        self.assertEqual(solicitud.sha_base_enviado, SHA_BASE_SIMULADO)
        self.assertEqual(solicitud.payload.get("sha"), SHA_BASE_SIMULADO)

        # Si falta el SHA al actualizar, debe rechazarse con ErrorShaBaseFaltanteGitHub
        with self.assertRaises(ErrorShaBaseFaltanteGitHub):
            construir_solicitud_put_contents_github(
                self.inventario_base,
                modo_publicacion=ModoPublicacionGitHub.ACTUALIZACION_ARCHIVO_EXISTENTE,
                sha_base_remoto=None,
                token=TOKEN_ESCRITURA_SIMULADO,
            )

    # ------------------------------------------------------------------
    # 11. Creación inicial permitida únicamente cuando GitHub confirma inexistencia real
    # ------------------------------------------------------------------
    def test_11_creacion_inicial_permitida_solo_cuando_github_confirma_inexistencia_real(
        self,
    ) -> None:
        estado, _ = self._crear_sesion_y_espacio_admin(
            inventario=self.inventario_base,
            sha_base=None,
            origen_remoto=False,
            es_copia_demostrativa=False,
            inicializacion_oficial_autorizada=True,
        )
        repo_memoria = RepositorioRemotoSimuladoMemoria(
            inventario_inicial=None,
            sha_inicial=None,
        )
        resultado = publicar_cambios_administrativos_github(
            estado,
            confirmacion_explicita=True,
            fuente_secretos=self.secretos_validos,
            consultor_existencia=repo_memoria.consultar_existencia,
            ejecutor_http_put=repo_memoria.ejecutar_put,
            reloj=self.reloj,
        )
        self.assertTrue(resultado.exito)
        self.assertTrue(resultado.confirmado_en_github)
        self.assertEqual(
            resultado.modo_publicacion,
            ModoPublicacionGitHub.CREACION_INICIAL_ARCHIVO_NUEVO,
        )
        self.assertEqual(
            resultado.estado,
            EstadoResultadoPublicacionGitHub.CREACION_INICIAL_CONFIRMADA,
        )
        self.assertEqual(len(repo_memoria.historial_puts), 1)
        payload_enviado = repo_memoria.historial_puts[0]["payload"]
        self.assertNotIn("sha", payload_enviado)
        self.assertEqual(payload_enviado["branch"], "main")

    # ------------------------------------------------------------------
    # 12. Rechazo de creación inicial si el archivo remoto aparece antes de publicar
    # ------------------------------------------------------------------
    def test_12_rechazo_creacion_inicial_si_archivo_remoto_aparece_antes_de_publicar(
        self,
    ) -> None:
        estado, _ = self._crear_sesion_y_espacio_admin(
            inventario=self.inventario_base,
            sha_base=None,
            origen_remoto=False,
            es_copia_demostrativa=False,
            inicializacion_oficial_autorizada=True,
        )
        puts_ejecutados: list[Any] = []

        def _put_no_debe_llamarse(*args: Any, **kwargs: Any) -> RespuestaHTTPGitHub:
            puts_ejecutados.append(args)
            raise AssertionError("No debió ejecutarse PUT si el archivo remoto ya existe.")

        with self.assertRaises(ErrorInventarioRemotoYaExiste) as ctx:
            publicar_cambios_administrativos_github(
                estado,
                confirmacion_explicita=True,
                fuente_secretos=self.secretos_validos,
                consultor_existencia=lambda **_: _existencia_remota_existe(
                    sha_github=SHA_REMOTO_CONCURRENTE
                ),
                ejecutor_http_put=_put_no_debe_llamarse,
                reloj=self.reloj,
            )
        self.assertEqual(
            ctx.exception.codigo,
            CodigoErrorPersistenciaGitHub.ARCHIVO_REMOTO_YA_EXISTE,
        )
        self.assertEqual(ctx.exception.sha_remoto_actual, SHA_REMOTO_CONCURRENTE)
        self.assertEqual(len(puts_ejecutados), 0)
        conflicto = estado.get(CLAVE_CONFLICTO_SHA_ADMIN)
        self.assertIsInstance(conflicto, EstadoConflictoShaAdmin)

    # ------------------------------------------------------------------
    # 13. Rechazo de creación inicial cuando la consulta previa devuelve error de red, auth o servidor
    # ------------------------------------------------------------------
    def test_13_rechazo_creacion_inicial_cuando_consulta_previa_devuelve_error(
        self,
    ) -> None:
        estado, _ = self._crear_sesion_y_espacio_admin(
            inventario=self.inventario_base,
            sha_base=None,
            origen_remoto=False,
            es_copia_demostrativa=False,
            inicializacion_oficial_autorizada=True,
        )
        codigos_error_previos = [
            (CodigoErrorGitHub.CONEXION_RED_FALLIDA, ErrorRedOTimeoutEscrituraGitHub),
            (CodigoErrorGitHub.CREDENCIAL_INVALIDA, ErrorCredencialesEscrituraGitHub),
            (CodigoErrorGitHub.ERROR_SERVIDOR, ErrorServidorEscrituraGitHub),
        ]
        for cod_gh, clase_exc in codigos_error_previos:
            existencia_fallida = ResultadoExistenciaRemota(
                estado=EstadoExistenciaRemota.DESCONOCIDO_ERROR,
                propietario="JunyMSAI",
                repositorio="FerreteriaMS",
                rama="main",
                ruta_archivo="03_Aplicacion/inventario.xlsx",
                codigo_error=cod_gh,
                categoria_consulta=CategoriaResultadoConsulta.ERROR_CONSULTA_INVENTARIO,
                mensaje=f"Fallo simulado previo: {cod_gh.value}",
            )
            with self.assertRaises(clase_exc):
                publicar_cambios_administrativos_github(
                    estado,
                    confirmacion_explicita=True,
                    fuente_secretos=self.secretos_validos,
                    consultor_existencia=lambda **_: existencia_fallida,
                    ejecutor_http_put=lambda *a, **k: (_ for _ in ()).throw(
                        AssertionError("No debía llamarse PUT ante error previo")
                    ),
                    reloj=self.reloj,
                )

    # ------------------------------------------------------------------
    # 14. Actualización simulada exitosa cuando el SHA base coincide con el SHA remoto actual
    # ------------------------------------------------------------------
    def test_14_actualizacion_simulada_exitosa_cuando_sha_base_coincide_con_remoto(
        self,
    ) -> None:
        estado, _ = self._crear_sesion_y_espacio_admin()
        crear_producto_local(
            estado,
            codigo_producto="0099",
            nombre="Taladro Percutor 1/2 Pulgada",
            categoria_id="CAT01",
            descripcion="Taladro industrial de velocidad variable.",
            ubicaciones_ids=["U001"],
            ubicacion_principal_id="U001",
            reloj=self.reloj,
        )
        resultado = publicar_cambios_administrativos_github(
            estado,
            confirmacion_explicita=True,
            fuente_secretos=self.secretos_validos,
            consultor_existencia=lambda **_: _existencia_remota_existe(SHA_BASE_SIMULADO),
            ejecutor_http_put=lambda *a, **k: _respuesta_put_exitosa_simulada(
                nuevo_sha_archivo=SHA_NUEVO_SIMULADO,
                sha_commit=SHA_COMMIT_SIMULADO,
            ),
            reloj=self.reloj,
        )
        self.assertTrue(resultado.exito)
        self.assertTrue(resultado.confirmado_en_github)
        self.assertEqual(
            resultado.estado,
            EstadoResultadoPublicacionGitHub.PUBLICACION_CONFIRMADA,
        )
        self.assertEqual(resultado.sha_anterior, SHA_BASE_SIMULADO)
        self.assertEqual(resultado.nuevo_sha_archivo, SHA_NUEVO_SIMULADO)
        self.assertEqual(resultado.sha_commit, SHA_COMMIT_SIMULADO)
        self.assertEqual(resultado.codigos_productos_creados, ("0099",))

    # ------------------------------------------------------------------
    # 15. Actualización del sha local al nuevo sha devuelto por GitHub tras publicar
    # ------------------------------------------------------------------
    def test_15_actualizacion_sha_local_al_nuevo_sha_devuelto_por_github(self) -> None:
        estado, _ = self._crear_sesion_y_espacio_admin()
        crear_producto_local(
            estado,
            codigo_producto="P-SHA-15",
            nombre="Llave Stilson 14 Pulgadas",
            categoria_id="CAT01",
            ubicaciones_ids=["U001"],
            ubicacion_principal_id="U001",
            reloj=self.reloj,
        )
        publicar_cambios_administrativos_github(
            estado,
            confirmacion_explicita=True,
            fuente_secretos=self.secretos_validos,
            consultor_existencia=lambda **_: _existencia_remota_existe(SHA_BASE_SIMULADO),
            ejecutor_http_put=lambda *a, **k: _respuesta_put_exitosa_simulada(
                nuevo_sha_archivo=SHA_NUEVO_SIMULADO
            ),
            reloj=self.reloj,
        )
        espacio_actualizado: EspacioTrabajoAdminInventario = estado[
            CLAVE_ESPACIO_TRABAJO_ADMIN
        ]
        self.assertEqual(espacio_actualizado.sha_base_remoto, SHA_NUEVO_SIMULADO)
        self.assertTrue(espacio_actualizado.origen_remoto_confirmado)
        self.assertTrue(espacio_actualizado.confirmado_en_github)
        self.assertFalse(espacio_actualizado.es_copia_local_temporal)

    # ------------------------------------------------------------------
    # 16. Limpieza del estado de cambios pendientes tras confirmación exitosa de GitHub
    # ------------------------------------------------------------------
    def test_16_limpieza_estado_cambios_pendientes_tras_confirmacion_exitosa(
        self,
    ) -> None:
        estado, _ = self._crear_sesion_y_espacio_admin()
        crear_producto_local(
            estado,
            codigo_producto="P-CLEAN-16",
            nombre="Cinta Métrica Profesional 8m",
            categoria_id="CAT03",
            ubicaciones_ids=["A-01"],
            ubicacion_principal_id="A-01",
            reloj=self.reloj,
        )
        editar_producto_local(
            estado,
            codigo_producto="0001",
            nombre="Tornillo Drywall 6x1 Negro Fosfatado",
            reloj=self.reloj,
        )
        publicar_cambios_administrativos_github(
            estado,
            confirmacion_explicita=True,
            fuente_secretos=self.secretos_validos,
            consultor_existencia=lambda **_: _existencia_remota_existe(SHA_BASE_SIMULADO),
            ejecutor_http_put=lambda *a, **k: _respuesta_put_exitosa_simulada(),
            reloj=self.reloj,
        )
        espacio_post: EspacioTrabajoAdminInventario = estado[CLAVE_ESPACIO_TRABAJO_ADMIN]
        self.assertFalse(espacio_post.tiene_cambios_locales_pendientes)
        self.assertEqual(espacio_post.codigos_productos_creados, ())
        self.assertEqual(espacio_post.codigos_productos_editados, ())
        self.assertEqual(espacio_post.historial_operaciones_locales, ())
        self.assertTrue(
            inventarios_son_identicos(
                espacio_post.inventario_original, espacio_post.inventario_trabajo
            )
        )
        prod_confirmado = consultar_producto_administrativo_por_codigo(
            estado, "P-CLEAN-16", reloj=self.reloj
        )
        self.assertTrue(prod_confirmado["confirmado_en_github"])
        self.assertFalse(prod_confirmado["tiene_cambios_locales"])

    # ------------------------------------------------------------------
    # 17. Limpieza o descarte del borrador temporal tras confirmación exitosa de GitHub
    # ------------------------------------------------------------------
    def test_17_limpieza_descarte_borrador_temporal_tras_confirmacion_exitosa(
        self,
    ) -> None:
        estado, _ = self._crear_sesion_y_espacio_admin()
        crear_producto_local(
            estado,
            codigo_producto="P-BORR-17",
            nombre="Nivel Magnético Torpedo 9 Pulgadas",
            categoria_id="CAT03",
            ubicaciones_ids=["U001"],
            ubicacion_principal_id="U001",
            reloj=self.reloj,
        )
        guardar_borrador_formulario_admin(
            estado,
            {"codigo_producto": "P-TEMP", "nombre": "Borrador en curso"},
            reloj=self.reloj,
        )
        self.assertTrue(tiene_borrador_pendiente_en_sesion(estado, reloj=self.reloj))
        self.assertIsNotNone(estado.get(CLAVE_FORMULARIO_BORRADOR_ADMIN))

        publicar_cambios_administrativos_github(
            estado,
            confirmacion_explicita=True,
            fuente_secretos=self.secretos_validos,
            consultor_existencia=lambda **_: _existencia_remota_existe(SHA_BASE_SIMULADO),
            ejecutor_http_put=lambda *a, **k: _respuesta_put_exitosa_simulada(),
            reloj=self.reloj,
        )
        self.assertFalse(tiene_borrador_pendiente_en_sesion(estado, reloj=self.reloj))
        self.assertIsNone(obtener_borrador_admin(estado, reloj=self.reloj))
        self.assertIsNone(estado.get(CLAVE_FORMULARIO_BORRADOR_ADMIN))

    # ------------------------------------------------------------------
    # 18. Detección de conflicto SHA cuando el SHA remoto cambió antes de publicar
    # ------------------------------------------------------------------
    def test_18_deteccion_conflicto_sha_cuando_sha_remoto_cambio_antes_de_publicar(
        self,
    ) -> None:
        estado, _ = self._crear_sesion_y_espacio_admin()
        crear_producto_local(
            estado,
            codigo_producto="P-CONF-18",
            nombre="Destornillador Estrella #2",
            categoria_id="CAT01",
            ubicaciones_ids=["U001"],
            ubicacion_principal_id="U001",
            reloj=self.reloj,
        )
        puts_llamados: list[Any] = []

        def _put_prohibido(*args: Any, **kwargs: Any) -> RespuestaHTTPGitHub:
            puts_llamados.append(args)
            raise AssertionError("No debe enviarse PUT cuando el SHA previo ya cambió.")

        with self.assertRaises(ErrorConflictoVersionShaGitHub) as ctx:
            publicar_cambios_administrativos_github(
                estado,
                confirmacion_explicita=True,
                fuente_secretos=self.secretos_validos,
                consultor_existencia=lambda **_: _existencia_remota_existe(
                    SHA_REMOTO_CONCURRENTE
                ),
                ejecutor_http_put=_put_prohibido,
                reloj=self.reloj,
            )
        self.assertEqual(
            ctx.exception.codigo,
            CodigoErrorPersistenciaGitHub.CONFLICTO_SHA_REMOTO,
        )
        self.assertEqual(ctx.exception.sha_base_local, SHA_BASE_SIMULADO)
        self.assertEqual(ctx.exception.sha_remoto_actual, SHA_REMOTO_CONCURRENTE)
        self.assertEqual(len(puts_llamados), 0)

        conflicto = estado.get(CLAVE_CONFLICTO_SHA_ADMIN)
        self.assertIsInstance(conflicto, EstadoConflictoShaAdmin)
        assert isinstance(conflicto, EstadoConflictoShaAdmin)
        self.assertEqual(conflicto.detectado_en, "verificacion_previa")
        self.assertEqual(conflicto.sha_base_local, SHA_BASE_SIMULADO)
        self.assertEqual(conflicto.sha_remoto_detectado, SHA_REMOTO_CONCURRENTE)

    # ------------------------------------------------------------------
    # 19. Detección de conflicto SHA cuando GitHub responde 409 Conflict o 422 por SHA desactualizado
    # ------------------------------------------------------------------
    def test_19_deteccion_conflicto_sha_cuando_github_responde_409_o_422(self) -> None:
        estado, _ = self._crear_sesion_y_espacio_admin()
        crear_producto_local(
            estado,
            codigo_producto="P-CONF-19",
            nombre="Alicate de Corte Diagonal 6 Pulgadas",
            categoria_id="CAT01",
            ubicaciones_ids=["U001"],
            ubicacion_principal_id="U001",
            reloj=self.reloj,
        )
        for status_conflicto in (409, 422):
            with self.assertRaises(ErrorConflictoVersionShaGitHub) as ctx:
                publicar_cambios_administrativos_github(
                    estado,
                    confirmacion_explicita=True,
                    fuente_secretos=self.secretos_validos,
                    consultor_existencia=lambda **_: _existencia_remota_existe(
                        SHA_BASE_SIMULADO
                    ),
                    ejecutor_http_put=lambda *a, sc=status_conflicto, **k: _respuesta_http_json(
                        sc,
                        {"message": "sha does not match latest blob sha"},
                    ),
                    reloj=self.reloj,
                )
            self.assertEqual(ctx.exception.status_code, status_conflicto)
            conflicto = estado.get(CLAVE_CONFLICTO_SHA_ADMIN)
            self.assertIsInstance(conflicto, EstadoConflictoShaAdmin)
            assert isinstance(conflicto, EstadoConflictoShaAdmin)
            self.assertEqual(conflicto.detectado_en, "respuesta_put")

    # ------------------------------------------------------------------
    # 20. Conservación intacta de cambios pendientes y borrador cuando ocurre conflicto SHA
    # ------------------------------------------------------------------
    def test_20_conservacion_intacta_cambios_pendientes_y_borrador_en_conflicto_sha(
        self,
    ) -> None:
        estado, _ = self._crear_sesion_y_espacio_admin()
        crear_producto_local(
            estado,
            codigo_producto="0077",
            nombre="Escuadra de Carpintero 12 Pulgadas",
            categoria_id="CAT03",
            ubicaciones_ids=["A-01"],
            ubicacion_principal_id="A-01",
            reloj=self.reloj,
        )
        editar_producto_local(
            estado,
            codigo_producto="A-01",
            nombre="Martillo de Uña 16 oz Mango Antivibración",
            reloj=self.reloj,
        )
        guardar_borrador_formulario_admin(
            estado,
            {"codigo_producto": "0078", "nombre": "Borrador activo protegido"},
            reloj=self.reloj,
        )

        res_ui = app_ferms.ejecutar_publicacion_github_en_sesion(
            estado,
            confirmacion_explicita=True,
            fuente_secretos=self.secretos_validos,
            consultor_existencia=lambda **_: _existencia_remota_existe(
                SHA_REMOTO_CONCURRENTE
            ),
            reloj=self.reloj,
        )
        self.assertFalse(res_ui.exito)
        self.assertFalse(res_ui.confirmado_en_github)
        self.assertTrue(res_ui.cambios_locales_preservados)
        self.assertTrue(res_ui.permite_recargar_remoto)
        self.assertEqual(
            res_ui.estado,
            EstadoResultadoPublicacionGitHub.CONFLICTO_SHA_REMOTO,
        )

        espacio_en_conflicto: EspacioTrabajoAdminInventario = estado[
            CLAVE_ESPACIO_TRABAJO_ADMIN
        ]
        self.assertTrue(espacio_en_conflicto.tiene_cambios_locales_pendientes)
        self.assertEqual(espacio_en_conflicto.codigos_productos_creados, ("0077",))
        self.assertEqual(espacio_en_conflicto.codigos_productos_editados, ("A-01",))
        self.assertEqual(espacio_en_conflicto.sha_base_remoto, SHA_BASE_SIMULADO)
        self.assertTrue(tiene_borrador_pendiente_en_sesion(estado, reloj=self.reloj))
        self.assertIsNotNone(estado.get(CLAVE_FORMULARIO_BORRADOR_ADMIN))

    # ------------------------------------------------------------------
    # 21. Opción controlada para recargar la versión remota tras conflicto sin pérdida silenciosa
    # ------------------------------------------------------------------
    def test_21_opcion_controlada_para_recargar_version_remota_tras_conflicto(
        self,
    ) -> None:
        estado, _ = self._crear_sesion_y_espacio_admin()
        crear_producto_local(
            estado,
            codigo_producto="P-CONF-21",
            nombre="Producto local en conflicto",
            categoria_id="CAT01",
            ubicaciones_ids=["U001"],
            ubicacion_principal_id="U001",
            reloj=self.reloj,
        )
        app_ferms.ejecutar_publicacion_github_en_sesion(
            estado,
            confirmacion_explicita=True,
            fuente_secretos=self.secretos_validos,
            consultor_existencia=lambda **_: _existencia_remota_existe(
                SHA_REMOTO_CONCURRENTE
            ),
            reloj=self.reloj,
        )

        # 1. Sin confirmación explícita, NO debe recargar ni descartar los cambios locales
        with self.assertRaises(ErrorConfirmacionPublicacionRequerida):
            recargar_inventario_remoto_tras_conflicto(
                estado,
                confirmar_descarte_locales=False,
                reloj=self.reloj,
            )
        espacio_aun_intacto: EspacioTrabajoAdminInventario = estado[
            CLAVE_ESPACIO_TRABAJO_ADMIN
        ]
        self.assertIn("P-CONF-21", espacio_aun_intacto.codigos_productos_creados)

        # 2. Con confirmación explícita, recarga el inventario remoto vigente y su nuevo SHA
        hojas_remotas = self.inventario_base.obtener_hojas()
        hojas_remotas[HOJA_PRODUCTOS].loc[0, "Nombre"] = "Tornillo Drywall 6x1 Remoto Actualizado"
        inv_remoto_nuevo = construir_inventario_datos(hojas_remotas)
        repo_actualizado = RepositorioRemotoSimuladoMemoria(
            inv_remoto_nuevo,
            sha_inicial=SHA_REMOTO_CONCURRENTE,
        )
        res_recarga = recargar_inventario_remoto_tras_conflicto(
            estado,
            confirmar_descarte_locales=True,
            descargador_remoto=repo_actualizado.descargar_remoto,
            reloj=self.reloj,
        )
        self.assertEqual(res_recarga.sha_base_remoto, SHA_REMOTO_CONCURRENTE)
        self.assertIsNone(estado.get(CLAVE_CONFLICTO_SHA_ADMIN))
        espacio_recargado: EspacioTrabajoAdminInventario = estado[
            CLAVE_ESPACIO_TRABAJO_ADMIN
        ]
        self.assertEqual(espacio_recargado.sha_base_remoto, SHA_REMOTO_CONCURRENTE)
        self.assertFalse(espacio_recargado.tiene_cambios_locales_pendientes)
        self.assertTrue(
            inventarios_son_identicos(
                espacio_recargado.inventario_trabajo, inv_remoto_nuevo
            )
        )

    # ------------------------------------------------------------------
    # 22. Manejo de error 401 Unauthorized sin declarar éxito ni perder cambios locales
    # ------------------------------------------------------------------
    def test_22_manejo_error_401_unauthorized_sin_declarar_exito_ni_perder_cambios(
        self,
    ) -> None:
        estado, _ = self._crear_sesion_y_espacio_admin()
        crear_producto_local(
            estado,
            codigo_producto="P-ERR-401",
            nombre="Producto ante error 401",
            categoria_id="CAT01",
            ubicaciones_ids=["U001"],
            ubicacion_principal_id="U001",
            reloj=self.reloj,
        )
        res = app_ferms.ejecutar_publicacion_github_en_sesion(
            estado,
            confirmacion_explicita=True,
            fuente_secretos=self.secretos_validos,
            consultor_existencia=lambda **_: _existencia_remota_existe(SHA_BASE_SIMULADO),
            ejecutor_http_put=lambda *a, **k: _respuesta_http_json(
                401, {"message": "Bad credentials"}
            ),
            reloj=self.reloj,
        )
        self.assertFalse(res.exito)
        self.assertFalse(res.confirmado_en_github)
        self.assertEqual(
            res.codigo_error,
            CodigoErrorPersistenciaGitHub.CREDENCIAL_INVALIDA_O_SIN_PERMISO,
        )
        self.assertTrue(estado[CLAVE_ESPACIO_TRABAJO_ADMIN].tiene_cambios_locales_pendientes)
        self.assertTrue(tiene_borrador_pendiente_en_sesion(estado, reloj=self.reloj))

    # ------------------------------------------------------------------
    # 23. Manejo de error 403 Forbidden o permisos insuficientes sin declarar éxito
    # ------------------------------------------------------------------
    def test_23_manejo_error_403_forbidden_o_permisos_insuficientes_sin_declarar_exito(
        self,
    ) -> None:
        estado, _ = self._crear_sesion_y_espacio_admin()
        crear_producto_local(
            estado,
            codigo_producto="P-ERR-403",
            nombre="Producto ante error 403",
            categoria_id="CAT01",
            ubicaciones_ids=["U001"],
            ubicacion_principal_id="U001",
            reloj=self.reloj,
        )
        # 403 por token sin permisos de escritura
        with self.assertRaises(ErrorCredencialesEscrituraGitHub) as ctx_perm:
            publicar_cambios_administrativos_github(
                estado,
                confirmacion_explicita=True,
                fuente_secretos=self.secretos_validos,
                consultor_existencia=lambda **_: _existencia_remota_existe(
                    SHA_BASE_SIMULADO
                ),
                ejecutor_http_put=lambda *a, **k: _respuesta_http_json(
                    403,
                    {"message": "Resource not accessible by personal access token"},
                ),
                reloj=self.reloj,
            )
        self.assertEqual(
            ctx_perm.exception.codigo,
            CodigoErrorPersistenciaGitHub.CREDENCIAL_INVALIDA_O_SIN_PERMISO,
        )

        # 403 por bloqueo de acceso general al repositorio
        with self.assertRaises(ErrorCredencialesEscrituraGitHub) as ctx_forb:
            publicar_cambios_administrativos_github(
                estado,
                confirmacion_explicita=True,
                fuente_secretos=self.secretos_validos,
                consultor_existencia=lambda **_: _existencia_remota_existe(
                    SHA_BASE_SIMULADO
                ),
                ejecutor_http_put=lambda *a, **k: _respuesta_http_json(
                    403,
                    {"message": "Repository access blocked"},
                ),
                reloj=self.reloj,
            )
        self.assertEqual(
            ctx_forb.exception.codigo,
            CodigoErrorPersistenciaGitHub.ACCESO_DENEGADO_ESCRITURA,
        )
        self.assertTrue(estado[CLAVE_ESPACIO_TRABAJO_ADMIN].tiene_cambios_locales_pendientes)

    # ------------------------------------------------------------------
    # 24. Manejo de error 404 de repositorio o rama sin confundirlo con inicialización permitida
    # ------------------------------------------------------------------
    def test_24_manejo_error_404_repositorio_o_rama_sin_confundir_con_inicializacion(
        self,
    ) -> None:
        estado, _ = self._crear_sesion_y_espacio_admin()
        crear_producto_local(
            estado,
            codigo_producto="P-ERR-404",
            nombre="Producto ante error 404 de rama o repo",
            categoria_id="CAT01",
            ubicaciones_ids=["U001"],
            ubicacion_principal_id="U001",
            reloj=self.reloj,
        )
        with self.assertRaises(ErrorRepositorioORamaEscrituraGitHub) as ctx_put_404:
            publicar_cambios_administrativos_github(
                estado,
                confirmacion_explicita=True,
                fuente_secretos=self.secretos_validos,
                consultor_existencia=lambda **_: _existencia_remota_existe(
                    SHA_BASE_SIMULADO
                ),
                ejecutor_http_put=lambda *a, **k: _respuesta_http_json(
                    404,
                    {"message": "Branch main not found"},
                ),
                reloj=self.reloj,
            )
        self.assertEqual(
            ctx_put_404.exception.codigo,
            CodigoErrorPersistenciaGitHub.REPOSITORIO_O_RAMA_INEXISTENTE,
        )
        self.assertTrue(estado[CLAVE_ESPACIO_TRABAJO_ADMIN].tiene_cambios_locales_pendientes)

    # ------------------------------------------------------------------
    # 25. Manejo de error 429 o rate limit agotado sin declarar éxito
    # ------------------------------------------------------------------
    def test_25_manejo_error_429_o_rate_limit_agotado_sin_declarar_exito(self) -> None:
        estado, _ = self._crear_sesion_y_espacio_admin()
        crear_producto_local(
            estado,
            codigo_producto="P-ERR-429",
            nombre="Producto ante rate limit",
            categoria_id="CAT01",
            ubicaciones_ids=["U001"],
            ubicacion_principal_id="U001",
            reloj=self.reloj,
        )
        for resp_rl in (
            _respuesta_http_json(429, {"message": "Too Many Requests"}),
            _respuesta_http_json(
                403,
                {"message": "API rate limit exceeded"},
                encabezados={"X-RateLimit-Remaining": "0"},
            ),
        ):
            with self.assertRaises(ErrorLimiteSolicitudesEscrituraGitHub) as ctx_rl:
                publicar_cambios_administrativos_github(
                    estado,
                    confirmacion_explicita=True,
                    fuente_secretos=self.secretos_validos,
                    consultor_existencia=lambda **_: _existencia_remota_existe(
                        SHA_BASE_SIMULADO
                    ),
                    ejecutor_http_put=lambda *a, r=resp_rl, **k: r,
                    reloj=self.reloj,
                )
            self.assertTrue(ctx_rl.exception.es_reintentable)
            self.assertTrue(
                estado[CLAVE_ESPACIO_TRABAJO_ADMIN].tiene_cambios_locales_pendientes
            )

    # ------------------------------------------------------------------
    # 26. Manejo de error 500/502/503 sin declarar éxito
    # ------------------------------------------------------------------
    def test_26_manejo_error_500_502_503_sin_declarar_exito(self) -> None:
        estado, _ = self._crear_sesion_y_espacio_admin()
        crear_producto_local(
            estado,
            codigo_producto="P-ERR-5XX",
            nombre="Producto ante fallo 5xx",
            categoria_id="CAT01",
            ubicaciones_ids=["U001"],
            ubicacion_principal_id="U001",
            reloj=self.reloj,
        )
        for codigo_5xx in (500, 502, 503):
            with self.assertRaises(ErrorServidorEscrituraGitHub) as ctx_5xx:
                publicar_cambios_administrativos_github(
                    estado,
                    confirmacion_explicita=True,
                    fuente_secretos=self.secretos_validos,
                    consultor_existencia=lambda **_: _existencia_remota_existe(
                        SHA_BASE_SIMULADO
                    ),
                    ejecutor_http_put=lambda *a, c=codigo_5xx, **k: _respuesta_http_json(
                        c, {"message": "Server Error"}
                    ),
                    reloj=self.reloj,
                )
            self.assertEqual(ctx_5xx.exception.status_code, codigo_5xx)
            self.assertTrue(ctx_5xx.exception.es_reintentable)
            self.assertTrue(
                estado[CLAVE_ESPACIO_TRABAJO_ADMIN].tiene_cambios_locales_pendientes
            )

    # ------------------------------------------------------------------
    # 27. Manejo de timeout o error de red sin declarar éxito ni perder cambios locales
    # ------------------------------------------------------------------
    def test_27_manejo_timeout_o_error_red_sin_declarar_exito_ni_perder_cambios(
        self,
    ) -> None:
        estado, _ = self._crear_sesion_y_espacio_admin()
        crear_producto_local(
            estado,
            codigo_producto="P-ERR-NET",
            nombre="Producto ante error de red o timeout",
            categoria_id="CAT01",
            ubicaciones_ids=["U001"],
            ubicacion_principal_id="U001",
            reloj=self.reloj,
        )
        # 1. Timeout
        with self.assertRaises(ErrorRedOTimeoutEscrituraGitHub) as ctx_to:
            publicar_cambios_administrativos_github(
                estado,
                confirmacion_explicita=True,
                fuente_secretos=self.secretos_validos,
                consultor_existencia=lambda **_: _existencia_remota_existe(
                    SHA_BASE_SIMULADO
                ),
                ejecutor_http_put=lambda *a, **k: (_ for _ in ()).throw(
                    socket.timeout("timed out writing to GitHub")
                ),
                reloj=self.reloj,
            )
        self.assertEqual(
            ctx_to.exception.codigo,
            CodigoErrorPersistenciaGitHub.TIEMPO_ESPERA_AGOTADO,
        )

        # 2. URLError de conexión
        with self.assertRaises(ErrorRedOTimeoutEscrituraGitHub) as ctx_net:
            publicar_cambios_administrativos_github(
                estado,
                confirmacion_explicita=True,
                fuente_secretos=self.secretos_validos,
                consultor_existencia=lambda **_: _existencia_remota_existe(
                    SHA_BASE_SIMULADO
                ),
                ejecutor_http_put=lambda *a, **k: (_ for _ in ()).throw(
                    urllib.error.URLError("Network unreachable")
                ),
                reloj=self.reloj,
            )
        self.assertEqual(
            ctx_net.exception.codigo,
            CodigoErrorPersistenciaGitHub.CONEXION_RED_FALLIDA,
        )
        self.assertTrue(estado[CLAVE_ESPACIO_TRABAJO_ADMIN].tiene_cambios_locales_pendientes)
        self.assertTrue(tiene_borrador_pendiente_en_sesion(estado, reloj=self.reloj))

    # ------------------------------------------------------------------
    # 28. Rechazo de respuesta 200/201 malformada o sin nuevo sha
    # ------------------------------------------------------------------
    def test_28_rechazo_respuesta_200_201_malformada_o_sin_nuevo_sha(self) -> None:
        estado, _ = self._crear_sesion_y_espacio_admin()
        crear_producto_local(
            estado,
            codigo_producto="P-ERR-200",
            nombre="Producto ante respuesta 200 sin SHA",
            categoria_id="CAT01",
            ubicaciones_ids=["U001"],
            ubicacion_principal_id="U001",
            reloj=self.reloj,
        )
        respuestas_invalidas = [
            RespuestaHTTPGitHub(status_code=200, cuerpo=b"No es JSON"),
            _respuesta_http_json(200, {"content": {}, "commit": {"sha": "123"}}),
            _respuesta_http_json(201, {"content": {"sha": ""}, "commit": {"sha": "123"}}),
            _respuesta_http_json(200, {"content": {"sha": "abc"}, "commit": {}}),
        ]
        for resp_inv in respuestas_invalidas:
            with self.assertRaises(ErrorRespuestaInesperadaEscrituraGitHub):
                publicar_cambios_administrativos_github(
                    estado,
                    confirmacion_explicita=True,
                    fuente_secretos=self.secretos_validos,
                    consultor_existencia=lambda **_: _existencia_remota_existe(
                        SHA_BASE_SIMULADO
                    ),
                    ejecutor_http_put=lambda *a, r=resp_inv, **k: r,
                    reloj=self.reloj,
                )
        self.assertTrue(estado[CLAVE_ESPACIO_TRABAJO_ADMIN].tiene_cambios_locales_pendientes)
        self.assertFalse(estado[CLAVE_ESPACIO_TRABAJO_ADMIN].confirmado_en_github)

    # ------------------------------------------------------------------
    # 29. Bloqueo de publicación sin sesión administrativa autenticada
    # ------------------------------------------------------------------
    def test_29_bloqueo_publicacion_sin_sesion_administrativa_autenticada(self) -> None:
        estado_no_auth: dict[str, Any] = {}
        app_ferms.inicializar_estado_sesion(estado_no_auth)
        with self.assertRaises(ErrorAutorizacionAdmin) as ctx:
            publicar_cambios_administrativos_github(
                estado_no_auth,
                confirmacion_explicita=True,
                fuente_secretos=self.secretos_validos,
                reloj=self.reloj,
            )
        self.assertEqual(
            ctx.exception.codigo,
            CodigoResultadoAutenticacion.CREDENCIALES_INVALIDAS,
        )

    # ------------------------------------------------------------------
    # 30. Bloqueo de publicación con sesión administrativa expirada y preservación del borrador
    # ------------------------------------------------------------------
    def test_30_bloqueo_publicacion_con_sesion_expirada_y_preservacion_borrador(
        self,
    ) -> None:
        estado, _ = self._crear_sesion_y_espacio_admin()
        crear_producto_local(
            estado,
            codigo_producto="P-EXP-30",
            nombre="Producto retenido en borrador al expirar",
            categoria_id="CAT01",
            ubicaciones_ids=["U001"],
            ubicacion_principal_id="U001",
            reloj=self.reloj,
        )
        # Avanzar más de 15 minutos de inactividad pero dentro de la ventana de 15 min adicionales del borrador
        self.reloj.avanzar(TIEMPO_EXPIRACION_INACTIVIDAD_SEGUNDOS + 10)

        res_exp = app_ferms.ejecutar_publicacion_github_en_sesion(
            estado,
            confirmacion_explicita=True,
            fuente_secretos=self.secretos_validos,
            reloj=self.reloj,
        )
        self.assertFalse(res_exp.exito)
        self.assertEqual(
            res_exp.estado,
            EstadoResultadoPublicacionGitHub.ERROR_AUTORIZACION_ADMIN,
        )
        self.assertTrue(tiene_borrador_pendiente_en_sesion(estado, reloj=self.reloj))

        # Reautenticar, restaurar el borrador y publicar exitosamente
        iniciar_sesion_admin(
            usuario=self.USUARIO_PRUEBA,
            contrasena=self.CONTRASENA_PRUEBA,
            estado=estado,
            fuente_secretos=self.secretos_validos,
            contexto_dispositivo=ContextoDispositivo.ESCRITORIO,
            reloj=self.reloj,
        )
        restaurar_borrador_en_espacio_trabajo_admin(
            estado,
            inventario_base=self.inventario_base,
            sha_base_remoto=SHA_BASE_SIMULADO,
            reloj=self.reloj,
        )
        res_post_rest = publicar_cambios_administrativos_github(
            estado,
            confirmacion_explicita=True,
            fuente_secretos=self.secretos_validos,
            consultor_existencia=lambda **_: _existencia_remota_existe(SHA_BASE_SIMULADO),
            ejecutor_http_put=lambda *a, **k: _respuesta_put_exitosa_simulada(),
            reloj=self.reloj,
        )
        self.assertTrue(res_post_rest.exito)
        self.assertIn("P-EXP-30", res_post_rest.codigos_productos_creados)

    # ------------------------------------------------------------------
    # 31. Bloqueo de publicación durante bloqueo administrativo activo
    # ------------------------------------------------------------------
    def test_31_bloqueo_publicacion_durante_bloqueo_administrativo_activo(self) -> None:
        estado, _ = self._crear_sesion_y_espacio_admin()
        crear_producto_local(
            estado,
            codigo_producto="P-LOCK-31",
            nombre="Producto bloqueado por intentos fallidos",
            categoria_id="CAT01",
            ubicaciones_ids=["U001"],
            ubicacion_principal_id="U001",
            reloj=self.reloj,
        )
        # Provocar bloqueo administrativo de 5 minutos mediante 5 intentos fallidos en otra sesión
        estado_otro: dict[str, Any] = {}
        for _ in range(5):
            iniciar_sesion_admin(
                usuario=self.USUARIO_PRUEBA,
                contrasena="ClaveIncorrecta!",
                estado=estado_otro,
                fuente_secretos=self.secretos_validos,
                contexto_dispositivo=ContextoDispositivo.ESCRITORIO,
                reloj=self.reloj,
            )

        with self.assertRaises(ErrorAutorizacionAdmin) as ctx_lock:
            publicar_cambios_administrativos_github(
                estado,
                confirmacion_explicita=True,
                fuente_secretos=self.secretos_validos,
                reloj=self.reloj,
            )
        self.assertEqual(
            ctx_lock.exception.codigo,
            CodigoResultadoAutenticacion.CUENTA_BLOQUEADA_TEMPORALMENTE,
        )

    # ------------------------------------------------------------------
    # 32. Bloqueo de publicación desde contexto móvil
    # ------------------------------------------------------------------
    def test_32_bloqueo_publicacion_desde_contexto_movil(self) -> None:
        estado, _ = self._crear_sesion_y_espacio_admin()
        crear_producto_local(
            estado,
            codigo_producto="P-MOVIL-32",
            nombre="Producto desde móvil",
            categoria_id="CAT01",
            ubicaciones_ids=["U001"],
            ubicacion_principal_id="U001",
            reloj=self.reloj,
        )
        with self.assertRaises(ErrorAutorizacionAdmin) as ctx_movil:
            publicar_cambios_administrativos_github(
                estado,
                confirmacion_explicita=True,
                fuente_secretos=self.secretos_validos,
                contexto_dispositivo=ContextoDispositivo.MOVIL,
                reloj=self.reloj,
            )
        self.assertEqual(
            ctx_movil.exception.codigo,
            CodigoResultadoAutenticacion.DISPOSITIVO_NO_PERMITIDO,
        )

    # ------------------------------------------------------------------
    # 33. Bloqueo de publicación sin confirmación explícita del administrador
    # ------------------------------------------------------------------
    def test_33_bloqueo_publicacion_sin_confirmacion_explicita_del_administrador(
        self,
    ) -> None:
        estado, _ = self._crear_sesion_y_espacio_admin()
        crear_producto_local(
            estado,
            codigo_producto="P-NOCONF-33",
            nombre="Producto sin confirmación explícita",
            categoria_id="CAT01",
            ubicaciones_ids=["U001"],
            ubicacion_principal_id="U001",
            reloj=self.reloj,
        )
        with self.assertRaises(ErrorConfirmacionPublicacionRequerida) as ctx:
            publicar_cambios_administrativos_github(
                estado,
                confirmacion_explicita=False,
                fuente_secretos=self.secretos_validos,
                reloj=self.reloj,
            )
        self.assertEqual(
            ctx.exception.codigo,
            CodigoErrorPersistenciaGitHub.CONFIRMACION_EXPLICITA_REQUERIDA,
        )
        self.assertTrue(estado[CLAVE_ESPACIO_TRABAJO_ADMIN].tiene_cambios_locales_pendientes)

    # ------------------------------------------------------------------
    # 34. Prevención de doble envío o publicación duplicada de la misma operación
    # ------------------------------------------------------------------
    def test_34_prevencion_doble_envio_o_publicacion_duplicada(self) -> None:
        estado, _ = self._crear_sesion_y_espacio_admin()
        crear_producto_local(
            estado,
            codigo_producto="P-DUP-34",
            nombre="Producto único contra doble envío",
            categoria_id="CAT01",
            ubicaciones_ids=["U001"],
            ubicacion_principal_id="U001",
            reloj=self.reloj,
        )
        repo_mem = RepositorioRemotoSimuladoMemoria(self.inventario_base)
        res_1 = publicar_cambios_administrativos_github(
            estado,
            confirmacion_explicita=True,
            fuente_secretos=self.secretos_validos,
            consultor_existencia=repo_mem.consultar_existencia,
            ejecutor_http_put=repo_mem.ejecutar_put,
            reloj=self.reloj,
        )
        self.assertTrue(res_1.exito)
        self.assertEqual(len(repo_mem.historial_puts), 1)

        # Segundo intento inmediato de publicar la misma operación confirmada
        with self.assertRaises(ErrorPublicacionDuplicadaGitHub) as ctx_dup:
            publicar_cambios_administrativos_github(
                estado,
                confirmacion_explicita=True,
                fuente_secretos=self.secretos_validos,
                consultor_existencia=repo_mem.consultar_existencia,
                ejecutor_http_put=repo_mem.ejecutar_put,
                reloj=self.reloj,
            )
        self.assertEqual(
            ctx_dup.exception.codigo,
            CodigoErrorPersistenciaGitHub.OPERACION_DUPLICADA,
        )
        self.assertEqual(len(repo_mem.historial_puts), 1)

    # ------------------------------------------------------------------
    # 35. Flujo completo: login -> crear producto -> confirmar -> publicar -> búsqueda pública
    # ------------------------------------------------------------------
    def test_35_flujo_completo_login_crear_producto_publicar_y_busqueda_publica(
        self,
    ) -> None:
        repo_mem = RepositorioRemotoSimuladoMemoria(self.inventario_base)
        estado: dict[str, Any] = {}
        app_ferms.inicializar_estado_sesion(estado)

        # 1. Login administrativo y carga desde el repositorio simulado
        res_login = app_ferms.ejecutar_login_administrativo_en_sesion(
            self.USUARIO_PRUEBA,
            self.CONTRASENA_PRUEBA,
            estado=estado,
            fuente_secretos=self.secretos_validos,
            contexto_dispositivo=ContextoDispositivo.ESCRITORIO,
            reloj=self.reloj,
        )
        self.assertTrue(res_login.autenticado)

        res_carga = app_ferms.ejecutar_carga_inventario_admin_en_sesion(
            estado,
            descargador_remoto=repo_mem.descargar_remoto,
            contexto_dispositivo=ContextoDispositivo.ESCRITORIO,
            reloj=self.reloj,
        )
        self.assertIsNotNone(res_carga)

        # 2. Crear producto nuevo con ceros a la izquierda
        res_creacion = app_ferms.ejecutar_creacion_producto_en_sesion(
            codigo_producto="0088",
            nombre="Esmeriladora Angular 4-1/2 Pulgadas",
            categoria_id="CAT01",
            descripcion="Esmeriladora compacta de 850W.",
            ubicaciones_ids=["U001", "A-01"],
            ubicacion_principal_id="A-01",
            estado=estado,
            contexto_dispositivo=ContextoDispositivo.ESCRITORIO,
            reloj=self.reloj,
        )
        self.assertTrue(res_creacion.exito)

        # 3. Confirmar y publicar en el repositorio simulado
        res_pub = app_ferms.ejecutar_publicacion_github_en_sesion(
            estado,
            confirmacion_explicita=True,
            fuente_secretos=self.secretos_validos,
            consultor_existencia=repo_mem.consultar_existencia,
            ejecutor_http_put=repo_mem.ejecutar_put,
            contexto_dispositivo=ContextoDispositivo.ESCRITORIO,
            reloj=self.reloj,
        )
        self.assertTrue(res_pub.exito)
        self.assertTrue(res_pub.confirmado_en_github)

        # 4. Nueva búsqueda pública encuentra inmediatamente el producto persistido
        resp_busqueda = app_ferms.ejecutar_flujo_busqueda_en_sesion(
            "Esmeriladora Angular",
            estado=estado,
            descargador_remoto=repo_mem.descargar_remoto,
        )
        self.assertEqual(
            resp_busqueda.estado_operativo,
            app_ferms.EstadoOperativoInterfaz.RESULTADOS_ENCONTRADOS,
        )
        assert resp_busqueda.resultado_busqueda is not None
        self.assertEqual(resp_busqueda.resultado_busqueda.total_resultados, 1)
        prod_encontrado = resp_busqueda.resultado_busqueda.resultados_pagina[0]
        self.assertEqual(prod_encontrado.codigo_producto, "0088")

        detalle = app_ferms.seleccionar_producto_en_sesion("0088", estado=estado)
        self.assertTrue(detalle.tiene_detalle)
        assert detalle.detalle is not None
        self.assertEqual(detalle.detalle.ubicacion_principal.ubicacion_id, "A-01")
        self.assertEqual(len(detalle.detalle.ubicaciones_secundarias), 1)
        self.assertEqual(detalle.detalle.ubicaciones_secundarias[0].ubicacion_id, "U001")

    # ------------------------------------------------------------------
    # 36. Flujo completo: login -> editar producto y ubicaciones -> publicar -> búsqueda pública
    # ------------------------------------------------------------------
    def test_36_flujo_completo_login_editar_producto_ubicaciones_publicar_y_busqueda_publica(
        self,
    ) -> None:
        repo_mem = RepositorioRemotoSimuladoMemoria(self.inventario_base)
        estado, _ = self._crear_sesion_y_espacio_admin()

        # Editar nombre, descripción y cambiar ubicación principal de '0001'
        res_ed = app_ferms.ejecutar_edicion_producto_en_sesion(
            codigo_producto="0001",
            nombre="Tornillo Drywall Reforzado 6x1",
            categoria_id="01",
            descripcion="Tornillo negro reforzado para placa de yeso.",
            ubicaciones_ids=["U001", "01", "U004"],
            ubicacion_principal_id="U004",
            estado=estado,
            contexto_dispositivo=ContextoDispositivo.ESCRITORIO,
            reloj=self.reloj,
        )
        self.assertTrue(res_ed.exito)

        res_pub = app_ferms.ejecutar_publicacion_github_en_sesion(
            estado,
            confirmacion_explicita=True,
            fuente_secretos=self.secretos_validos,
            consultor_existencia=repo_mem.consultar_existencia,
            ejecutor_http_put=repo_mem.ejecutar_put,
            contexto_dispositivo=ContextoDispositivo.ESCRITORIO,
            reloj=self.reloj,
        )
        self.assertTrue(res_pub.exito)

        resp_busqueda = app_ferms.ejecutar_flujo_busqueda_en_sesion(
            "Tornillo Drywall Reforzado",
            estado=estado,
            descargador_remoto=repo_mem.descargar_remoto,
        )
        self.assertEqual(
            resp_busqueda.estado_operativo,
            app_ferms.EstadoOperativoInterfaz.RESULTADOS_ENCONTRADOS,
        )
        detalle = app_ferms.seleccionar_producto_en_sesion("0001", estado=estado)
        assert detalle.detalle is not None
        self.assertEqual(detalle.detalle.nombre, "Tornillo Drywall Reforzado 6x1")
        self.assertEqual(detalle.detalle.ubicacion_principal.ubicacion_id, "U004")
        ids_secundarias = [
            u.ubicacion_id for u in detalle.detalle.ubicaciones_secundarias
        ]
        self.assertEqual(sorted(ids_secundarias), ["01", "U001"])

    # ------------------------------------------------------------------
    # 37. Verificación de que cambios locales no confirmados no aparecen en búsqueda pública
    # ------------------------------------------------------------------
    def test_37_cambios_locales_no_confirmados_no_aparecen_en_busqueda_publica(
        self,
    ) -> None:
        repo_mem = RepositorioRemotoSimuladoMemoria(self.inventario_base)
        estado, _ = self._crear_sesion_y_espacio_admin()

        crear_producto_local(
            estado,
            codigo_producto="P-OCULTO-37",
            nombre="Cepillo Carpintero Exclusivo Local",
            categoria_id="CAT01",
            ubicaciones_ids=["U001"],
            ubicacion_principal_id="U001",
            reloj=self.reloj,
        )

        # La búsqueda pública consulta GitHub (donde aún no se ha publicado) y NO debe ver P-OCULTO-37
        resp_busqueda = app_ferms.ejecutar_flujo_busqueda_en_sesion(
            "Cepillo Carpintero Exclusivo Local",
            estado=estado,
            descargador_remoto=repo_mem.descargar_remoto,
        )
        self.assertEqual(
            resp_busqueda.estado_operativo,
            app_ferms.EstadoOperativoInterfaz.SIN_COINCIDENCIAS,
        )

    # ------------------------------------------------------------------
    # 38. Verificación de que no se modifica inventario.xlsx local ni se ejecutan solicitudes reales
    # ------------------------------------------------------------------
    def test_38_no_se_crea_ni_modifica_inventario_xlsx_local_ni_solicitudes_reales(
        self,
    ) -> None:
        self.assertTrue(
            RUTA_INVENTARIO_XLSX.is_file(),
            "El archivo oficial 03_Aplicacion/inventario.xlsx debe existir.",
        )
        self.assertEqual(
            RUTA_INVENTARIO_XLSX.read_bytes(),
            BYTES_INVENTARIO_XLSX_INICIAL,
            "El archivo 03_Aplicacion/inventario.xlsx no debe modificarse localmente.",
        )
        self.assertFalse(
            ESCRITURA_REAL_GITHUB_HABILITADA,
            "La compuerta de escritura real a GitHub debe permanecer deshabilitada por defecto.",
        )
        # Intentar invocar ejecutar_solicitud_put_github sin mock de urlopen ni autorización real debe bloquearse
        solicitud, _ = construir_solicitud_put_contents_github(
            self.inventario_base,
            modo_publicacion=ModoPublicacionGitHub.ACTUALIZACION_ARCHIVO_EXISTENTE,
            sha_base_remoto=SHA_BASE_SIMULADO,
            token=TOKEN_ESCRITURA_SIMULADO,
        )
        with self.assertRaises(ErrorEscrituraRealNoAutorizadaGitHub):
            ejecutar_solicitud_put_github(
                solicitud.url,
                solicitud.encabezados,
                solicitud.cuerpo_json_bytes,
            )
        self.assertEqual(RUTA_INVENTARIO_XLSX.read_bytes(), BYTES_INVENTARIO_XLSX_INICIAL)
        inv_oficial = cargar_inventario_excel(RUTA_INVENTARIO_XLSX)
        self.assertEqual(tuple(inv_oficial.obtener_hojas().keys()), ORDEN_HOJAS)

    # ------------------------------------------------------------------
    # 39. Ausencia de exposición de secretos, tokens, contraseñas o hashes
    # ------------------------------------------------------------------
    def test_39_ausencia_exposicion_secretos_tokens_contrasenas_hashes(self) -> None:
        # 1. Verificar login real con Argon2id sin mocks
        self.patcher_argon2.stop()
        try:
            estado_argon: dict[str, Any] = {}
            app_ferms.inicializar_estado_sesion(estado_argon)
            res_auth = iniciar_sesion_admin(
                usuario=self.USUARIO_PRUEBA,
                contrasena=self.CONTRASENA_PRUEBA,
                estado=estado_argon,
                fuente_secretos=self.secretos_validos,
                contexto_dispositivo=ContextoDispositivo.ESCRITORIO,
                reloj=self.reloj,
            )
            self.assertTrue(res_auth.autenticado)
        finally:
            self.patcher_argon2.start()

        # 2. Verificar sanitización en excepciones de persistencia
        token_secreto = "ghp_99887766554433221100aabbccddeeffGGHH"
        err = ErrorCredencialesEscrituraGitHub(
            f"Fallo con token {token_secreto} y hash {self.HASH_PRUEBA_ARGON2ID}",
            detalle_tecnico=f"Bearer {token_secreto} / password={self.CONTRASENA_PRUEBA}",
            token=token_secreto,
        )
        self.assertNotIn(token_secreto, str(err))
        self.assertNotIn(token_secreto, err.mensaje)
        self.assertNotIn(token_secreto, err.detalle_tecnico)
        self.assertNotIn("$argon2id$", err.mensaje)
        self.assertNotIn(self.CONTRASENA_PRUEBA, repr(estado_argon))
        self.assertNotIn(self.HASH_PRUEBA_ARGON2ID, repr(estado_argon))
        self.assertNotIn(TOKEN_ESCRITURA_SIMULADO, repr(estado_argon))

    # ------------------------------------------------------------------
    # 40. Integración Streamlit AppTest P10 y regresión de contratos P02..P09
    # ------------------------------------------------------------------
    def test_40_integracion_interfaz_streamlit_p10_y_regresion_prompts_02_a_09(
        self,
    ) -> None:
        repo_mem = RepositorioRemotoSimuladoMemoria(self.inventario_base)

        with patch.dict(
            os.environ, self.secretos_validos, clear=False
        ), patch.object(
            gh,
            "descargar_inventario_remoto",
            side_effect=repo_mem.descargar_remoto,
        ), patch.object(
            persist_gh,
            "consultar_existencia_inventario_remoto",
            side_effect=repo_mem.consultar_existencia,
        ), patch.object(
            persist_gh,
            "ejecutar_solicitud_put_github",
            side_effect=repo_mem.ejecutar_put,
        ):
            at = AppTest.from_file(str(RUTA_APP), default_timeout=15)
            at.run()
            self.assertFalse(at.exception)

            # 1. Navegar a Productos e iniciar sesión
            at.button(key="btn_nav_productos").click()
            at.run()
            at.text_input(key="campo_admin_usuario").set_value(self.USUARIO_PRUEBA)
            at.text_input(key="campo_admin_contrasena").set_value(
                self.CONTRASENA_PRUEBA
            )
            at.button(key="btn_admin_iniciar_sesion").click()
            at.run()
            self.assertFalse(at.exception)

            # 2. Crear un producto local
            at.button(key="btn_admin_abrir_crear").click()
            at.run()
            at.text_input(key="campo_admin_crear_codigo").set_value("0055")
            at.text_input(key="campo_admin_crear_nombre").set_value(
                "Serrucho Costilla 12 Pulgadas"
            )
            at.text_area(key="campo_admin_crear_descripcion").set_value(
                "Serrucho de precisión para cortes de inglete."
            )
            at.button(key="btn_admin_confirmar_crear_producto").click()
            at.run()
            self.assertFalse(at.exception)

            # 3. Intentar publicar sin marcar el checkbox de confirmación explícita
            at.button(key="btn_admin_publicar_github").click()
            at.run()
            self.assertFalse(at.exception)
            warnings_ui = [w.value for w in at.warning]
            self.assertTrue(
                any(
                    MENSAJE_CONFIRMACION_EXPLICITA_REQUERIDA in w
                    for w in warnings_ui
                )
            )
            self.assertEqual(len(repo_mem.historial_puts), 0)

            # 4. Marcar confirmación explícita y publicar en el repositorio simulado
            at.checkbox(key="chk_admin_confirmar_publicacion_github").check()
            at.run()
            at.button(key="btn_admin_publicar_github").click()
            at.run()
            self.assertFalse(at.exception)
            self.assertEqual(len(repo_mem.historial_puts), 1)
            exitos_ui = [s.value for s in at.success]
            self.assertTrue(
                any(
                    MENSAJE_PUBLICACION_ACTUALIZACION_EXITOSA in s
                    for s in exitos_ui
                )
            )

            # 5. Volver al buscador público y buscar el producto recién publicado sin reiniciar Streamlit
            at.button(key="btn_nav_inicio").click()
            at.run()
            at.text_input(key="campo_nombre_producto").set_value("Serrucho Costilla")
            at.button(key="btn_ejecutar_busqueda").click()
            at.run()
            self.assertFalse(at.exception)

            textos_markdown = " ".join(m.value for m in at.markdown)
            self.assertIn("Serrucho Costilla 12 Pulgadas", textos_markdown)
            self.assertIn("0055", textos_markdown)

    # ------------------------------------------------------------------
    # 41. Corrección 1 de auditoría: Rechazo de intentos de anular la compuerta de escritura real
    # ------------------------------------------------------------------
    def test_41_rechazo_intentos_anular_compuerta_escritura_real_github(self) -> None:
        # 1. Ninguna función pública acepta el parámetro 'permitir_red_real'
        for func_publica in (
            ejecutar_solicitud_put_github,
            publicar_cambios_administrativos_github,
            app_ferms.ejecutar_publicacion_github_en_sesion,
        ):
            parametros = inspect.signature(func_publica).parameters
            self.assertNotIn(
                "permitir_red_real",
                parametros,
                f"La función {func_publica.__name__} no debe exponer el parámetro 'permitir_red_real'.",
            )

        solicitud, _ = construir_solicitud_put_contents_github(
            self.inventario_base,
            modo_publicacion=ModoPublicacionGitHub.ACTUALIZACION_ARCHIVO_EXISTENTE,
            sha_base_remoto=SHA_BASE_SIMULADO,
            token=TOKEN_ESCRITURA_SIMULADO,
        )
        with self.assertRaises(TypeError):
            ejecutar_solicitud_put_github(  # type: ignore[call-arg]
                solicitud.url,
                solicitud.encabezados,
                solicitud.cuerpo_json_bytes,
                permitir_red_real=True,
            )

        # 2. Sustituir urllib.request.urlopen (incluso con assert_called o __module__ distinto)
        # NO permite que ejecutar_solicitud_put_github salte ESCRITURA_REAL_GITHUB_HABILITADA=False
        with patch("urllib.request.urlopen") as mock_urlopen:
            with self.assertRaises(ErrorEscrituraRealNoAutorizadaGitHub) as ctx_gate:
                ejecutar_solicitud_put_github(
                    solicitud.url,
                    solicitud.encabezados,
                    solicitud.cuerpo_json_bytes,
                )
            self.assertEqual(
                ctx_gate.exception.codigo,
                CodigoErrorPersistenciaGitHub.ESCRITURA_REAL_NO_AUTORIZADA,
            )
            mock_urlopen.assert_not_called()

        # 3. En publicar_cambios_administrativos_github, omitir ejecutor_http_put o inyectar
        # funciones de red reales es rechazado con ErrorEscrituraRealNoAutorizadaGitHub
        estado, _ = self._crear_sesion_y_espacio_admin()
        crear_producto_local(
            estado,
            codigo_producto="P-GATE-41",
            nombre="Producto bloqueado por compuerta global",
            categoria_id="CAT01",
            ubicaciones_ids=["U001"],
            ubicacion_principal_id="U001",
            reloj=self.reloj,
        )
        for ejecutor_prohibido in (
            None,
            ejecutar_solicitud_put_github,
            urllib.request.urlopen,
            urllib.request.build_opener().open,
        ):
            with self.assertRaises(ErrorEscrituraRealNoAutorizadaGitHub):
                publicar_cambios_administrativos_github(
                    estado,
                    confirmacion_explicita=True,
                    fuente_secretos=self.secretos_validos,
                    consultor_existencia=lambda **_: _existencia_remota_existe(
                        SHA_BASE_SIMULADO
                    ),
                    ejecutor_http_put=ejecutor_prohibido,  # type: ignore[arg-type]
                    reloj=self.reloj,
                )

        # 4. Ante un resultado de escritura ambiguo (p. ej. respuesta 200 sin SHA o timeout),
        # se revoca la confirmación en sesión y no se permite reutilizar la misma instancia de solicitud
        estado[CLAVE_CONFIRMACION_GUARDADO_GITHUB] = True
        with self.assertRaises(ErrorRespuestaInesperadaEscrituraGitHub):
            publicar_cambios_administrativos_github(
                estado,
                confirmacion_explicita=True,
                fuente_secretos=self.secretos_validos,
                consultor_existencia=lambda **_: _existencia_remota_existe(
                    SHA_BASE_SIMULADO
                ),
                ejecutor_http_put=lambda *a, **k: RespuestaHTTPGitHub(
                    status_code=200, cuerpo=b"{}"
                ),
                reloj=self.reloj,
            )
        self.assertFalse(estado[CLAVE_CONFIRMACION_GUARDADO_GITHUB])
        persist_gh._registrar_despacho_unico_solicitud_put(solicitud)
        self.assertTrue(getattr(solicitud, "_despachada", False))
        with self.assertRaises(ErrorPublicacionDuplicadaGitHub):
            persist_gh._registrar_despacho_unico_solicitud_put(solicitud)
        self.assertTrue(estado[CLAVE_ESPACIO_TRABAJO_ADMIN].tiene_cambios_locales_pendientes)

    # ------------------------------------------------------------------
    # 42. Corrección 2 de auditoría: Bloqueo de publicación de datos demostrativos y separación oficial
    # ------------------------------------------------------------------
    def test_42_bloqueo_publicacion_datos_demostrativos_y_separacion_inicializacion_oficial(
        self,
    ) -> None:
        # 1. construir_solicitud_put_contents_github rechaza inventarios demostrativos
        inv_demo = crear_inventario_demostrativo()
        with self.assertRaises(ErrorPublicacionDatosDemostrativosNoAutorizada) as ctx_demo_req:
            construir_solicitud_put_contents_github(
                inv_demo,
                modo_publicacion=ModoPublicacionGitHub.CREACION_INICIAL_ARCHIVO_NUEVO,
                token=TOKEN_ESCRITURA_SIMULADO,
            )
        self.assertEqual(
            ctx_demo_req.exception.codigo,
            CodigoErrorPersistenciaGitHub.DATOS_DEMOSTRATIVOS_NO_PUBLICABLES,
        )

        # 2. inicializar_espacio_trabajo_desde_inventario rechaza marcar como oficial un inventario demo
        estado_rechazo: dict[str, Any] = {}
        app_ferms.inicializar_estado_sesion(estado_rechazo)
        iniciar_sesion_admin(
            usuario=self.USUARIO_PRUEBA,
            contrasena=self.CONTRASENA_PRUEBA,
            estado=estado_rechazo,
            fuente_secretos=self.secretos_validos,
            contexto_dispositivo=ContextoDispositivo.ESCRITORIO,
            reloj=self.reloj,
        )
        with self.assertRaises(ErrorValidacionAdministrativa):
            inicializar_espacio_trabajo_desde_inventario(
                estado_rechazo,
                inv_demo,
                sha_base_remoto=None,
                origen_remoto_confirmado=False,
                inicializacion_oficial_autorizada=True,
                contexto_dispositivo=ContextoDispositivo.ESCRITORIO,
                reloj=self.reloj,
            )

        # 3. Una copia demostrativa preparada en sesión cuando no existe remoto no puede publicarse
        def _lector_inexistente(**_kwargs: Any) -> ResultadoLecturaInventarioRemoto:
            raise gh.ErrorArchivoRemotoInexistente("No existe inventario.xlsx.")

        res_carga_demo = cargar_inventario_administrativo_en_sesion(
            estado_rechazo,
            descargador_remoto=_lector_inexistente,
            inicializar_demo_si_no_existe=True,
            contexto_dispositivo=ContextoDispositivo.ESCRITORIO,
            reloj=self.reloj,
        )
        assert res_carga_demo.espacio_trabajo is not None
        self.assertTrue(res_carga_demo.espacio_trabajo.es_copia_demostrativa)
        self.assertFalse(res_carga_demo.espacio_trabajo.inicializacion_oficial_autorizada)

        puts_intentados: list[Any] = []
        res_pub_demo = app_ferms.ejecutar_publicacion_github_en_sesion(
            estado_rechazo,
            confirmacion_explicita=True,
            fuente_secretos=self.secretos_validos,
            consultor_existencia=lambda **_: _existencia_remota_no_existe_confirmada(),
            ejecutor_http_put=lambda *a, **k: (
                puts_intentados.append(a) or _respuesta_put_exitosa_simulada()
            ),
            contexto_dispositivo=ContextoDispositivo.ESCRITORIO,
            reloj=self.reloj,
        )
        self.assertFalse(res_pub_demo.exito)
        self.assertFalse(res_pub_demo.confirmado_en_github)
        self.assertEqual(
            res_pub_demo.estado,
            EstadoResultadoPublicacionGitHub.BLOQUEADO_DATOS_DEMOSTRATIVOS,
        )
        self.assertEqual(
            res_pub_demo.codigo_error,
            CodigoErrorPersistenciaGitHub.DATOS_DEMOSTRATIVOS_NO_PUBLICABLES,
        )
        self.assertEqual(len(puts_intentados), 0)

        # 4. Incluso con datos no demostrativos, una creación inicial sin inicializacion_oficial_autorizada=True es rechazada
        estado_no_aut, _ = self._crear_sesion_y_espacio_admin(
            inventario=self.inventario_base,
            sha_base=None,
            origen_remoto=False,
            es_copia_demostrativa=False,
            inicializacion_oficial_autorizada=False,
        )
        with self.assertRaises(ErrorPublicacionDatosDemostrativosNoAutorizada) as ctx_init_no_aut:
            publicar_cambios_administrativos_github(
                estado_no_aut,
                confirmacion_explicita=True,
                fuente_secretos=self.secretos_validos,
                consultor_existencia=lambda **_: _existencia_remota_no_existe_confirmada(),
                ejecutor_http_put=lambda *a, **k: _respuesta_put_exitosa_simulada(),
                reloj=self.reloj,
            )
        self.assertEqual(
            ctx_init_no_aut.exception.codigo,
            CodigoErrorPersistenciaGitHub.INICIALIZACION_OFICIAL_NO_AUTORIZADA,
        )

        # 5. En la interfaz Streamlit, al preparar la copia demostrativa en memoria NO se ofrece inicializar en GitHub
        with patch.dict(
            os.environ, self.secretos_validos, clear=False
        ), patch.object(
            gh,
            "descargar_inventario_remoto",
            side_effect=_lector_inexistente,
        ):
            at = AppTest.from_file(str(RUTA_APP), default_timeout=15)
            at.run()
            at.button(key="btn_nav_productos").click()
            at.run()
            at.text_input(key="campo_admin_usuario").set_value(self.USUARIO_PRUEBA)
            at.text_input(key="campo_admin_contrasena").set_value(self.CONTRASENA_PRUEBA)
            at.button(key="btn_admin_iniciar_sesion").click()
            at.run()
            at.button(key="btn_admin_cargar_inventario").click()
            at.run()
            at.button(key="btn_admin_inicializar_demo_local").click()
            at.run()
            self.assertFalse(at.exception)

            claves_botones = [b.key for b in at.button]
            claves_checkbox = [c.key for c in at.checkbox]
            self.assertNotIn("btn_admin_publicar_github", claves_botones)
            self.assertNotIn("chk_admin_confirmar_publicacion_github", claves_checkbox)
            textos_info = " ".join(i.value for i in at.info)
            self.assertIn("[DEMOSTRATIVO]", textos_info)

    # ------------------------------------------------------------------
    # 43. Verificación integral de secretos en repr/str/logs y archivos versionados
    # ------------------------------------------------------------------
    def test_43_auditoria_estatica_y_dinamica_no_exposicion_secretos_y_gitignore(
        self,
    ) -> None:
        token_ficticio = "ghp_TokenFicticioAuditoriaSecretaP10_ABCDEF123456"
        token_custom = "tok_secreto_personalizado_auditoria_987654321"
        clave_ficticia = "ClaveFicticiaSuperSecreta!2026"
        hash_ficticio = self.HASH_PRUEBA_ARGON2ID

        solicitud, paquete = construir_solicitud_put_contents_github(
            self.inventario_base,
            modo_publicacion=ModoPublicacionGitHub.ACTUALIZACION_ARCHIVO_EXISTENTE,
            sha_base_remoto=SHA_BASE_SIMULADO,
            token=token_ficticio,
        )
        conflicto = EstadoConflictoShaAdmin(
            sha_base_local=SHA_BASE_SIMULADO,
            sha_remoto_detectado=SHA_REMOTO_CONCURRENTE,
            modo_publicacion=ModoPublicacionGitHub.ACTUALIZACION_ARCHIVO_EXISTENTE,
            detectado_en="verificacion_previa",
            mensaje_usuario=f"Conflicto con Authorization: Bearer {token_custom}",
            detalle_tecnico=f"Detalle con token {token_ficticio} y hash {hash_ficticio}",
        )
        resultado = persist_gh.ResultadoPublicacionGitHub(
            exito=False,
            confirmado_en_github=False,
            estado=EstadoResultadoPublicacionGitHub.ERROR_CREDENCIALES_O_PERMISOS,
            modo_publicacion=ModoPublicacionGitHub.ACTUALIZACION_ARCHIVO_EXISTENTE,
            mensaje_usuario=f"Error con password={clave_ficticia}",
            detalle_tecnico=f"Authorization: Bearer {token_ficticio}",
            codigo_error=CodigoErrorPersistenciaGitHub.CREDENCIAL_INVALIDA_O_SIN_PERMISO,
        )
        encabezado_dict_repr = repr({"Authorization": f"Bearer {token_custom}"})
        excepcion = ErrorPersistenciaGitHub(
            f"Excepción con Bearer {token_ficticio} y password={clave_ficticia}",
            detalle_tecnico=f"Hash {hash_ficticio} y {encabezado_dict_repr}",
            token=token_ficticio,
        )

        for obj in (solicitud, paquete, conflicto, resultado, excepcion):
            for representacion in (str(obj), repr(obj)):
                self.assertNotIn(token_ficticio, representacion)
                self.assertNotIn(token_custom, representacion)
                self.assertNotIn(clave_ficticia, representacion)
                self.assertNotIn(hash_ficticio, representacion)
                self.assertNotIn("$argon2id$", representacion)

        # Verificar exclusiones de secretos en .gitignore
        ruta_gitignore = RAIZ_PROYECTO / ".gitignore"
        self.assertTrue(ruta_gitignore.is_file())
        contenido_gitignore = ruta_gitignore.read_text(encoding="utf-8")
        self.assertIn(".streamlit/secrets.toml", contenido_gitignore)
        self.assertIn(".env", contenido_gitignore)
        self.assertIn("*.pem", contenido_gitignore)
        self.assertIn("*.key", contenido_gitignore)
        self.assertIn("~$*.xlsx", contenido_gitignore)

        # Verificar archivos versionados mediante git ls-files
        proc = subprocess.run(
            ["git", "ls-files"],
            cwd=str(RAIZ_PROYECTO),
            capture_output=True,
            text=True,
            check=True,
        )
        archivos_versionados = [
            linea.strip() for linea in proc.stdout.splitlines() if linea.strip()
        ]
        self.assertIn("03_Aplicacion/inventario.xlsx", archivos_versionados)
        self.assertEqual(RUTA_INVENTARIO_XLSX.read_bytes(), BYTES_INVENTARIO_XLSX_INICIAL)
        self.assertFalse(
            any(a.endswith("secrets.toml") for a in archivos_versionados),
            "Ningún archivo secrets.toml debe estar versionado en git.",
        )

        # Escaneo estático de módulos de producción en 03_Aplicacion/
        patron_secreto_real = re.compile(
            r"(ghp_[A-Za-z0-9]{30,}|github_pat_[A-Za-z0-9_]{30,}|AIza[0-9A-Za-z\-_]{30,}|\$argon2id\$v=19\$m=\d+,t=\d+,p=\d+\$[A-Za-z0-9+/=]+\$[A-Za-z0-9+/=]+)"
        )
        for archivo_py in RUTA_APLICACION.glob("*.py"):
            texto_py = archivo_py.read_text(encoding="utf-8")
            self.assertIsNone(
                patron_secreto_real.search(texto_py),
                f"Se detectó un patrón de credencial o hash literal en {archivo_py.name}.",
            )


if __name__ == "__main__":
    unittest.main(verbosity=2)
