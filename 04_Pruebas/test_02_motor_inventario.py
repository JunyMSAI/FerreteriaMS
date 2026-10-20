"""Pruebas unitarias y de integración local para el motor del inventario Excel.

Proyecto: FerreteríaMS v0
Tarea: FERMS-V0-P02
Módulo probado: 03_Aplicacion/inventario.py
Requisitos cubiertos: RF-020, RF-022, RF-023, RF-024, RF-029 (parcial), RF-030 (parcial), RF-031.
Criterios cubiertos: CA-020 (parcial), CA-021 (parcial), CA-022, CA-047, CA-048, CA-049, CA-050.

Todas las pruebas utilizan memoria o directorios temporales aislados (tempfile.TemporaryDirectory)
y verifican que no se cree ni modifique 03_Aplicacion/inventario.xlsx.
"""

from __future__ import annotations

import hashlib
import io
import sys
import tempfile
import unittest
from pathlib import Path

import openpyxl
import pandas as pd

RAIZ_PROYECTO = Path(__file__).resolve().parent.parent
RUTA_APLICACION = RAIZ_PROYECTO / "03_Aplicacion"
RUTA_MAESTRO_DEFINITIVO = RUTA_APLICACION / "inventario.xlsx"

if str(RUTA_APLICACION) not in sys.path:
    sys.path.insert(0, str(RUTA_APLICACION))

import inventario as motor  # noqa: E402


def _crear_hojas_validas_base() -> dict[str, pd.DataFrame]:
    """Construye un conjunto base de cuatro hojas válidas con datos ficticios de prueba."""
    return {
        motor.HOJA_CATEGORIAS: pd.DataFrame(
            [
                {"Categoria_ID": "01", "Nombre_Categoria": "Herramientas (Demostrativo)"},
                {"Categoria_ID": "02", "Nombre_Categoria": "Tornillería (Demostrativo)"},
            ]
        ),
        motor.HOJA_UBICACIONES: pd.DataFrame(
            [
                {"Ubicacion_ID": "001", "Rack": "R01", "Columna": "C01", "Nivel": "N01", "Posicion": "P01"},
                {"Ubicacion_ID": "002", "Rack": "R02", "Columna": "C03", "Nivel": "N01", "Posicion": "P02"},
                {"Ubicacion_ID": "003", "Rack": "R03", "Columna": "C02", "Nivel": "N04", "Posicion": "P01"},
            ]
        ),
        motor.HOJA_PRODUCTOS: pd.DataFrame(
            [
                {
                    "Codigo_Producto": "0007",
                    "Nombre": "Martillo de uña",
                    "Categoria_ID": "01",
                    "Descripcion": "[DEMOSTRATIVO] Martillo de acero forjado",
                },
                {
                    "Codigo_Producto": "TOR001",
                    "Nombre": "Tornillo galvanizado",
                    "Categoria_ID": "02",
                    "Descripcion": "",
                },
            ]
        ),
        motor.HOJA_ASIGNACIONES: pd.DataFrame(
            [
                {"Codigo_Producto": "0007", "Ubicacion_ID": "001", "Principal": "Sí"},
                {"Codigo_Producto": "0007", "Ubicacion_ID": "002", "Principal": "No"},
                {"Codigo_Producto": "TOR001", "Ubicacion_ID": "003", "Principal": "si"},
            ]
        ),
    }


def _escribir_hojas_en_xlsx(hojas: dict[str, pd.DataFrame], ruta: Path) -> None:
    """Escribe un diccionario de DataFrames en un archivo .xlsx temporal."""
    wb = openpyxl.Workbook()
    hoja_inicial = wb.active
    if hoja_inicial is not None:
        wb.remove(hoja_inicial)

    for nombre_hoja, df in hojas.items():
        ws = wb.create_sheet(title=nombre_hoja)
        columnas = list(df.columns)
        for col_idx, col_name in enumerate(columnas, start=1):
            celda = ws.cell(row=1, column=col_idx, value=str(col_name))
            celda.number_format = "@"
            celda.data_type = "s"
        for row_idx, (_, fila) in enumerate(df.iterrows(), start=2):
            for col_idx, col_name in enumerate(columnas, start=1):
                val = "" if pd.isna(fila[col_name]) else str(fila[col_name])
                celda = ws.cell(row=row_idx, column=col_idx, value=val)
                celda.number_format = "@"
                celda.data_type = "s"
    wb.save(ruta)


class TestMotorInventarioCasosValidos(unittest.TestCase):
    """Pruebas de la Etapa D — 7.1 Casos válidos."""

    def test_lectura_cuatro_hojas_y_encabezados(self) -> None:
        """CA-022 / RF-023: Reconoce las cuatro hojas y sus encabezados exactos."""
        with tempfile.TemporaryDirectory() as tmpdir:
            ruta_temp = Path(tmpdir) / "inventario_valido.xlsx"
            hojas = _crear_hojas_validas_base()
            _escribir_hojas_en_xlsx(hojas, ruta_temp)

            inv = motor.cargar_inventario_excel(ruta_temp)
            hojas_cargadas = inv.obtener_hojas()

            self.assertEqual(set(hojas_cargadas.keys()), set(motor.ESQUEMA_HOJAS.keys()))
            for nombre_hoja, cols_esperadas in motor.ESQUEMA_HOJAS.items():
                self.assertEqual(list(hojas_cargadas[nombre_hoja].columns), cols_esperadas)

    def test_conservacion_identificadores_como_texto_con_ceros_iniciales(self) -> None:
        """RF-023 / RF-024: Conserva códigos e identificadores como texto incluyendo ceros iniciales."""
        with tempfile.TemporaryDirectory() as tmpdir:
            ruta_temp = Path(tmpdir) / "ceros_iniciales.xlsx"
            hojas = _crear_hojas_validas_base()
            _escribir_hojas_en_xlsx(hojas, ruta_temp)

            inv = motor.cargar_inventario_excel(ruta_temp)
            self.assertIn("0007", list(inv.productos["Codigo_Producto"]))
            self.assertIn("01", list(inv.categorias["Categoria_ID"]))
            self.assertIn("001", list(inv.ubicaciones["Ubicacion_ID"]))
            for col in inv.productos.columns:
                for valor in inv.productos[col]:
                    self.assertIsInstance(valor, str)

    def test_productos_categorias_y_ubicaciones_principal_y_secundarias(self) -> None:
        """RF-006 / RF-007 / RF-019: Estructura correctamente ubicación principal y secundarias."""
        hojas = _crear_hojas_validas_base()
        inv = motor.construir_inventario_datos(hojas)

        prod_multiple = inv.obtener_producto_por_codigo("0007")
        self.assertIsNotNone(prod_multiple)
        assert prod_multiple is not None
        self.assertEqual(prod_multiple["nombre"], "Martillo de uña")
        self.assertEqual(prod_multiple["categoria_id"], "01")
        self.assertEqual(prod_multiple["nombre_categoria"], "Herramientas (Demostrativo)")
        self.assertIsNotNone(prod_multiple["ubicacion_principal"])
        self.assertEqual(prod_multiple["ubicacion_principal"]["ubicacion_id"], "001")
        self.assertEqual(prod_multiple["ubicacion_principal"]["rack"], "R01")
        self.assertEqual(len(prod_multiple["ubicaciones_secundarias"]), 1)
        self.assertEqual(prod_multiple["ubicaciones_secundarias"][0]["ubicacion_id"], "002")

        prod_unico = inv.obtener_producto_por_codigo("TOR001")
        self.assertIsNotNone(prod_unico)
        assert prod_unico is not None
        self.assertEqual(prod_unico["descripcion"], "")
        self.assertIsNotNone(prod_unico["ubicacion_principal"])
        self.assertEqual(prod_unico["ubicacion_principal"]["ubicacion_id"], "003")
        self.assertEqual(prod_unico["ubicaciones_secundarias"], [])

        tabla = inv.construir_tabla_consulta()
        self.assertEqual(len(tabla), 3)

    def test_generacion_demostrativa_en_memoria_y_serializacion_ciclica(self) -> None:
        """RF-029 / RF-031 / CA-020 parcial: Genera inventario demostrativo válido y serializable."""
        inv_demo = motor.crear_inventario_demostrativo()
        resultado_val = motor.validar_inventario(inv_demo.obtener_hojas(), lanzar_excepcion=False)
        self.assertTrue(resultado_val.es_valido)

        # Verificar marca demostrativa RF-031
        descripciones = list(inv_demo.productos["Descripcion"])
        self.assertTrue(all(motor.MARCA_DEMOSTRATIVA in d for d in descripciones))
        self.assertIn("000125", list(inv_demo.productos["Codigo_Producto"]))

        with tempfile.TemporaryDirectory() as tmpdir:
            ruta_temp = Path(tmpdir) / "demo_ciclo.xlsx"
            motor.guardar_inventario_en_ruta(inv_demo, ruta_temp)
            inv_releido = motor.cargar_inventario_excel(ruta_temp)

            self.assertEqual(
                list(inv_releido.productos["Codigo_Producto"]),
                list(inv_demo.productos["Codigo_Producto"]),
            )
            self.assertFalse(RUTA_MAESTRO_DEFINITIVO.exists())


class TestMotorInventarioCasosInvalidos(unittest.TestCase):
    """Pruebas de la Etapa D — 7.2 Casos inválidos."""

    def test_producto_con_codigo_duplicado(self) -> None:
        """RF-020 / CA-014: Rechaza códigos de producto duplicados."""
        hojas = _crear_hojas_validas_base()
        hojas[motor.HOJA_PRODUCTOS] = pd.concat(
            [
                hojas[motor.HOJA_PRODUCTOS],
                pd.DataFrame(
                    [
                        {
                            "Codigo_Producto": "0007",
                            "Nombre": "Otro producto con mismo código",
                            "Categoria_ID": "01",
                            "Descripcion": "",
                        }
                    ]
                ),
            ],
            ignore_index=True,
        )
        with self.assertRaises(motor.ErrorIdentificadorDuplicado) as ctx:
            motor.validar_inventario(hojas, lanzar_excepcion=True)
        self.assertEqual(ctx.exception.codigo, motor.CodigoErrorInventario.IDENTIFICADOR_DUPLICADO)

    def test_producto_sin_nombre_o_sin_codigo(self) -> None:
        """RF-020: Rechaza productos con campos obligatorios vacíos (Nombre o Codigo_Producto)."""
        hojas = _crear_hojas_validas_base()
        hojas[motor.HOJA_PRODUCTOS].loc[0, "Nombre"] = "   "
        with self.assertRaises(motor.ErrorValorObligatorio) as ctx:
            motor.validar_inventario(hojas, lanzar_excepcion=True)
        self.assertEqual(ctx.exception.codigo, motor.CodigoErrorInventario.VALOR_OBLIGATORIO_AUSENTE)

    def test_producto_con_categoria_inexistente(self) -> None:
        """RF-020 / CA-047: Rechaza productos asociados a una categoría que no existe."""
        hojas = _crear_hojas_validas_base()
        hojas[motor.HOJA_PRODUCTOS].loc[0, "Categoria_ID"] = "CAT_FANTASMA"
        with self.assertRaises(motor.ErrorCategoriaInexistente) as ctx:
            motor.validar_inventario(hojas, lanzar_excepcion=True)
        self.assertEqual(ctx.exception.codigo, motor.CodigoErrorInventario.CATEGORIA_INEXISTENTE)

    def test_categorias_duplicadas_por_identificador(self) -> None:
        """RF-023: Rechaza categorías con Categoria_ID duplicado."""
        hojas = _crear_hojas_validas_base()
        hojas[motor.HOJA_CATEGORIAS] = pd.concat(
            [
                hojas[motor.HOJA_CATEGORIAS],
                pd.DataFrame([{"Categoria_ID": "01", "Nombre_Categoria": "Categoría Repetida"}]),
            ],
            ignore_index=True,
        )
        with self.assertRaises(motor.ErrorIdentificadorDuplicado) as ctx:
            motor.validar_inventario(hojas, lanzar_excepcion=True)
        self.assertEqual(ctx.exception.codigo, motor.CodigoErrorInventario.IDENTIFICADOR_DUPLICADO)

    def test_ubicaciones_duplicadas_por_identificador(self) -> None:
        """RF-023: Rechaza ubicaciones con Ubicacion_ID duplicado."""
        hojas = _crear_hojas_validas_base()
        hojas[motor.HOJA_UBICACIONES] = pd.concat(
            [
                hojas[motor.HOJA_UBICACIONES],
                pd.DataFrame(
                    [
                        {
                            "Ubicacion_ID": "001",
                            "Rack": "R09",
                            "Columna": "C09",
                            "Nivel": "N09",
                            "Posicion": "P09",
                        }
                    ]
                ),
            ],
            ignore_index=True,
        )
        with self.assertRaises(motor.ErrorIdentificadorDuplicado) as ctx:
            motor.validar_inventario(hojas, lanzar_excepcion=True)
        self.assertEqual(ctx.exception.codigo, motor.CodigoErrorInventario.IDENTIFICADOR_DUPLICADO)

    def test_combinaciones_fisicas_duplicadas(self) -> None:
        """RF-023 / CA-049: Detecta y rechaza combinaciones Rack + Columna + Nivel + Posicion duplicadas."""
        hojas = _crear_hojas_validas_base()
        hojas[motor.HOJA_UBICACIONES] = pd.concat(
            [
                hojas[motor.HOJA_UBICACIONES],
                pd.DataFrame(
                    [
                        {
                            "Ubicacion_ID": "999",
                            "Rack": "R01",
                            "Columna": "C01",
                            "Nivel": "N01",
                            "Posicion": "P01",
                        }
                    ]
                ),
            ],
            ignore_index=True,
        )
        with self.assertRaises(motor.ErrorUbicacionFisicaDuplicada) as ctx:
            motor.validar_inventario(hojas, lanzar_excepcion=True)
        self.assertEqual(
            ctx.exception.codigo, motor.CodigoErrorInventario.UBICACION_FISICA_DUPLICADA
        )

    def test_asignacion_a_producto_inexistente(self) -> None:
        """RF-020 / RF-023: Rechaza asignaciones cuyo Codigo_Producto no existe en Productos."""
        hojas = _crear_hojas_validas_base()
        hojas[motor.HOJA_ASIGNACIONES] = pd.concat(
            [
                hojas[motor.HOJA_ASIGNACIONES],
                pd.DataFrame(
                    [
                        {
                            "Codigo_Producto": "PROD_NO_EXISTE",
                            "Ubicacion_ID": "001",
                            "Principal": "No",
                        }
                    ]
                ),
            ],
            ignore_index=True,
        )
        with self.assertRaises(motor.ErrorProductoInexistente) as ctx:
            motor.validar_inventario(hojas, lanzar_excepcion=True)
        self.assertEqual(ctx.exception.codigo, motor.CodigoErrorInventario.PRODUCTO_INEXISTENTE)

    def test_asignacion_a_ubicacion_inexistente(self) -> None:
        """RF-020 / CA-048: Rechaza asignaciones cuya Ubicacion_ID no existe en Ubicaciones."""
        hojas = _crear_hojas_validas_base()
        hojas[motor.HOJA_ASIGNACIONES] = pd.concat(
            [
                hojas[motor.HOJA_ASIGNACIONES],
                pd.DataFrame(
                    [
                        {
                            "Codigo_Producto": "TOR001",
                            "Ubicacion_ID": "UBIC_INEXISTENTE",
                            "Principal": "No",
                        }
                    ]
                ),
            ],
            ignore_index=True,
        )
        with self.assertRaises(motor.ErrorUbicacionInexistente) as ctx:
            motor.validar_inventario(hojas, lanzar_excepcion=True)
        self.assertEqual(ctx.exception.codigo, motor.CodigoErrorInventario.UBICACION_INEXISTENTE)

    def test_asignacion_duplicada_producto_ubicacion(self) -> None:
        """RF-023 / CA-050: Detecta y rechaza relaciones duplicadas entre el mismo producto y ubicación."""
        hojas = _crear_hojas_validas_base()
        hojas[motor.HOJA_ASIGNACIONES] = pd.concat(
            [
                hojas[motor.HOJA_ASIGNACIONES],
                pd.DataFrame(
                    [{"Codigo_Producto": "0007", "Ubicacion_ID": "002", "Principal": "No"}]
                ),
            ],
            ignore_index=True,
        )
        with self.assertRaises(motor.ErrorAsignacionDuplicada) as ctx:
            motor.validar_inventario(hojas, lanzar_excepcion=True)
        self.assertEqual(ctx.exception.codigo, motor.CodigoErrorInventario.ASIGNACION_DUPLICADA)

    def test_producto_sin_ubicacion_asignada(self) -> None:
        """RF-019 / RF-020: Rechaza productos que no tengan al menos una ubicación asignada."""
        hojas = _crear_hojas_validas_base()
        hojas[motor.HOJA_PRODUCTOS] = pd.concat(
            [
                hojas[motor.HOJA_PRODUCTOS],
                pd.DataFrame(
                    [
                        {
                            "Codigo_Producto": "SIN_UBIC_01",
                            "Nombre": "Alicate aislado",
                            "Categoria_ID": "01",
                            "Descripcion": "",
                        }
                    ]
                ),
            ],
            ignore_index=True,
        )
        with self.assertRaises(motor.ErrorProductoSinAsignacion) as ctx:
            motor.validar_inventario(hojas, lanzar_excepcion=True)
        self.assertEqual(ctx.exception.codigo, motor.CodigoErrorInventario.PRODUCTO_SIN_ASIGNACION)

    def test_producto_con_dos_ubicaciones_principales(self) -> None:
        """RF-019 / RF-020 / CA-046: Rechaza productos con más de una ubicación principal."""
        hojas = _crear_hojas_validas_base()
        # Marcar también como principal la segunda ubicación de '0007'
        hojas[motor.HOJA_ASIGNACIONES].loc[1, "Principal"] = "Sí"
        with self.assertRaises(motor.ErrorUbicacionPrincipalMultiple) as ctx:
            motor.validar_inventario(hojas, lanzar_excepcion=True)
        self.assertEqual(
            ctx.exception.codigo, motor.CodigoErrorInventario.UBICACION_PRINCIPAL_MULTIPLE
        )

    def test_producto_sin_ubicacion_principal(self) -> None:
        """RF-019 / RF-020: Rechaza productos cuyas asignaciones sean todas secundarias."""
        hojas = _crear_hojas_validas_base()
        hojas[motor.HOJA_ASIGNACIONES].loc[0, "Principal"] = "No"
        with self.assertRaises(motor.ErrorSinUbicacionPrincipal) as ctx:
            motor.validar_inventario(hojas, lanzar_excepcion=True)
        self.assertEqual(ctx.exception.codigo, motor.CodigoErrorInventario.SIN_UBICACION_PRINCIPAL)

    def test_valores_invalidos_del_campo_principal(self) -> None:
        """RF-020 / RF-023: Rechaza valores ambiguos o inválidos en el campo Principal."""
        for valor_invalido in ["Tal vez", "1", "True", "Principal", "", "S", "N"]:
            hojas = _crear_hojas_validas_base()
            hojas[motor.HOJA_ASIGNACIONES].loc[0, "Principal"] = valor_invalido
            with self.assertRaises(motor.ErrorValorPrincipalInvalido) as ctx:
                motor.validar_inventario(hojas, lanzar_excepcion=True)
            self.assertEqual(
                ctx.exception.codigo, motor.CodigoErrorInventario.VALOR_PRINCIPAL_INVALIDO
            )

    def test_archivo_con_hoja_o_columna_faltante_o_no_autorizada(self) -> None:
        """RF-023 / CA-022: Rechaza archivos con hojas faltantes, columnas faltantes o columnas comerciales."""
        # 1. Hoja faltante
        hojas_incompletas = _crear_hojas_validas_base()
        del hojas_incompletas[motor.HOJA_UBICACIONES]
        with self.assertRaises(motor.ErrorHojaFaltante) as ctx_hoja:
            motor.validar_inventario(hojas_incompletas, lanzar_excepcion=True)
        self.assertEqual(ctx_hoja.exception.codigo, motor.CodigoErrorInventario.HOJA_FALTANTE)

        # 2. Columna obligatoria faltante
        hojas_sin_col = _crear_hojas_validas_base()
        hojas_sin_col[motor.HOJA_PRODUCTOS] = hojas_sin_col[motor.HOJA_PRODUCTOS].drop(
            columns=["Nombre"]
        )
        with self.assertRaises(motor.ErrorColumnaFaltante) as ctx_col:
            motor.validar_inventario(hojas_sin_col, lanzar_excepcion=True)
        self.assertEqual(ctx_col.exception.codigo, motor.CodigoErrorInventario.COLUMNA_FALTANTE)

        # 3. Columna comercial no autorizada (ej. Precio)
        hojas_con_precio = _crear_hojas_validas_base()
        hojas_con_precio[motor.HOJA_PRODUCTOS]["Precio"] = "10.00"
        with self.assertRaises(motor.ErrorColumnaFaltante) as ctx_extra:
            motor.validar_inventario(hojas_con_precio, lanzar_excepcion=True)
        self.assertEqual(
            ctx_extra.exception.codigo, motor.CodigoErrorInventario.COLUMNA_NO_AUTORIZADA
        )

    def test_archivo_con_hoja_adicional_no_autorizada(self) -> None:
        """RF-023 / CA-022: Rechaza archivos o diccionarios que contengan hojas adicionales no autorizadas."""
        # 1. Validación directa de estructura con hoja adicional en memoria
        hojas_con_extra = _crear_hojas_validas_base()
        hojas_con_extra["Existencias"] = pd.DataFrame(
            [{"Codigo_Producto": "0007", "Cantidad": "15"}]
        )
        resultado = motor.validar_estructura_hojas(hojas_con_extra, lanzar_excepcion=False)
        self.assertFalse(resultado.es_valido)
        self.assertIn(motor.CodigoErrorInventario.HOJA_NO_AUTORIZADA, resultado.codigos_error)

        with self.assertRaises(motor.ErrorHojaNoAutorizada) as ctx_mem:
            motor.validar_estructura_hojas(hojas_con_extra, lanzar_excepcion=True)
        self.assertEqual(ctx_mem.exception.codigo, motor.CodigoErrorInventario.HOJA_NO_AUTORIZADA)

        # 2. Lectura de archivo Excel (.xlsx) temporal con una quinta hoja no autorizada
        with tempfile.TemporaryDirectory() as tmpdir:
            ruta_extra = Path(tmpdir) / "inventario_hoja_extra.xlsx"
            _escribir_hojas_en_xlsx(hojas_con_extra, ruta_extra)
            with self.assertRaises(motor.ErrorHojaNoAutorizada) as ctx_xlsx:
                motor.cargar_inventario_excel(ruta_extra)
            self.assertEqual(
                ctx_xlsx.exception.codigo, motor.CodigoErrorInventario.HOJA_NO_AUTORIZADA
            )

    def test_archivo_inexistente_y_archivo_corrupto(self) -> None:
        """Sección 5.4: Diferencia claramente entre archivo inexistente y archivo corrupto."""
        with tempfile.TemporaryDirectory() as tmpdir:
            ruta_inexistente = Path(tmpdir) / "no_existe.xlsx"
            with self.assertRaises(motor.ErrorArchivoInexistente) as ctx_no_existe:
                motor.cargar_inventario_excel(ruta_inexistente)
            self.assertEqual(
                ctx_no_existe.exception.codigo, motor.CodigoErrorInventario.ARCHIVO_INEXISTENTE
            )

            ruta_corrupta = Path(tmpdir) / "corrupto.xlsx"
            ruta_corrupta.write_bytes(b"ESTO_NO_ES_UN_ARCHIVO_EXCEL_VALIDO")
            with self.assertRaises(motor.ErrorArchivoCorrupto) as ctx_corrupto:
                motor.cargar_inventario_excel(ruta_corrupta)
            self.assertEqual(
                ctx_corrupto.exception.codigo, motor.CodigoErrorInventario.ARCHIVO_CORRUPTO
            )

            with self.assertRaises(motor.ErrorArchivoCorrupto):
                motor.cargar_inventario_excel(io.BytesIO(b""))


class TestProteccionDatosInventario(unittest.TestCase):
    """Pruebas de la Etapa D — 7.3 Protección de datos."""

    def test_validacion_no_sobrescribe_archivo_entrada(self) -> None:
        """Sección 5.4 / 7.3: Ni la carga válida ni la fallida alteran los bytes del archivo de entrada."""
        with tempfile.TemporaryDirectory() as tmpdir:
            # Archivo válido
            ruta_valida = Path(tmpdir) / "entrada_valida.xlsx"
            _escribir_hojas_en_xlsx(_crear_hojas_validas_base(), ruta_valida)
            hash_antes_valido = hashlib.sha256(ruta_valida.read_bytes()).hexdigest()

            motor.cargar_inventario_excel(ruta_valida)
            hash_despues_valido = hashlib.sha256(ruta_valida.read_bytes()).hexdigest()
            self.assertEqual(hash_antes_valido, hash_despues_valido)

            # Archivo con error de integridad
            hojas_invalidas = _crear_hojas_validas_base()
            hojas_invalidas[motor.HOJA_PRODUCTOS].loc[0, "Categoria_ID"] = "INEXISTENTE"
            ruta_invalida = Path(tmpdir) / "entrada_invalida.xlsx"
            _escribir_hojas_en_xlsx(hojas_invalidas, ruta_invalida)
            hash_antes_invalido = hashlib.sha256(ruta_invalida.read_bytes()).hexdigest()

            with self.assertRaises(motor.ErrorCategoriaInexistente):
                motor.cargar_inventario_excel(ruta_invalida)
            hash_despues_invalido = hashlib.sha256(ruta_invalida.read_bytes()).hexdigest()
            self.assertEqual(hash_antes_invalido, hash_despues_invalido)

    def test_proteccion_contra_sobrescritura_e_inicializacion_segura(self) -> None:
        """RF-029 / RF-030 / CA-020 parcial / CA-021 parcial: Protege archivos existentes y estados inciertos."""
        with tempfile.TemporaryDirectory() as tmpdir:
            ruta_existente = Path(tmpdir) / "inventario.xlsx"
            contenido_original = b"CONTENIDO_EXISTENTE_PROTEGIDO"
            ruta_existente.write_bytes(contenido_original)

            inv_demo = motor.crear_inventario_demostrativo()

            # 1. guardar_inventario_en_ruta debe rechazar sobrescribir por defecto
            with self.assertRaises(motor.ErrorProteccionInventario) as ctx_guardado:
                motor.guardar_inventario_en_ruta(inv_demo, ruta_existente)
            self.assertEqual(
                ctx_guardado.exception.codigo,
                motor.CodigoErrorInventario.PROTECCION_INVENTARIO_EXISTENTE,
            )
            self.assertEqual(ruta_existente.read_bytes(), contenido_original)

            # 2. Estado remoto EXISTE debe bloquear inicialización
            with self.assertRaises(motor.ErrorProteccionInventario) as ctx_remoto_existe:
                motor.preparar_inicializacion_segura(motor.EstadoExistenciaRemota.EXISTE)
            self.assertEqual(
                ctx_remoto_existe.exception.codigo,
                motor.CodigoErrorInventario.PROTECCION_INVENTARIO_EXISTENTE,
            )

            # 3. Estado remoto DESCONOCIDO_ERROR (ej. fallo de red/permisos) debe bloquear inicialización
            with self.assertRaises(motor.ErrorProteccionInventario) as ctx_incierto:
                motor.preparar_inicializacion_segura(motor.EstadoExistenciaRemota.DESCONOCIDO_ERROR)
            self.assertEqual(
                ctx_incierto.exception.codigo,
                motor.CodigoErrorInventario.ESTADO_REMOTO_DESCONOCIDO,
            )

            # 4. Estado remoto NO_EXISTE pero archivo local existente también debe bloquearse
            with self.assertRaises(motor.ErrorProteccionInventario):
                motor.preparar_inicializacion_segura(
                    motor.EstadoExistenciaRemota.NO_EXISTE,
                    ruta_local_control=ruta_existente,
                )
            self.assertEqual(ruta_existente.read_bytes(), contenido_original)

            # 5. Solo con confirmación explícita NO_EXISTE y sin archivo local existente se generan los bytes
            ruta_no_existente = Path(tmpdir) / "nuevo_maestro.xlsx"
            bytes_generados = motor.preparar_inicializacion_segura(
                motor.EstadoExistenciaRemota.NO_EXISTE,
                ruta_local_control=ruta_no_existente,
            )
            self.assertGreater(len(bytes_generados), 0)
            # Y aun así no se escribió automáticamente en disco
            self.assertFalse(ruta_no_existente.exists())

    def test_importar_modulo_no_crea_inventario_maestro(self) -> None:
        """Sección 6.2 / 7.3: Verifica que 03_Aplicacion/inventario.xlsx no haya sido creado."""
        self.assertFalse(
            RUTA_MAESTRO_DEFINITIVO.exists(),
            "El archivo maestro definitivo 03_Aplicacion/inventario.xlsx no debe crearse en el Prompt 02.",
        )


if __name__ == "__main__":
    unittest.main(verbosity=2)

